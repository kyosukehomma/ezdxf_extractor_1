import hashlib
import json
import re
import shutil
import tempfile
import unittest
from decimal import Decimal
from pathlib import Path
from unittest.mock import patch

import ezdxf
from openpyxl import load_workbook

import oudan_red_yellow_checker as app


PROJECT_ROOT = Path(__file__).resolve().parents[1]
LOCAL_CASE_DIR = PROJECT_ROOT / "input" / "used_260723_162452"
VERIFIED_OUTPUT_DIR = PROJECT_ROOT / "output" / "exec_260723_162452"

BAD_EXPECTED_STATS = {
    "オープン掘削": (17, Decimal("838.0"), 60),
    "路体盛土(W＜2.5m)": (3, Decimal("0.8"), 74),
    "路体盛土(2.5m≦W＜4.0m)": (2, Decimal("1.0"), 75),
    "路体盛土(4.0m≦W)": (56, Decimal("11944.3"), 21),
    "路床盛土(W＜2.5m)": (1, Decimal("0.2"), 76),
    "路床盛土(2.5m≦W＜4.0m)": (1, Decimal("0.7"), 76),
    "路床盛土(4.0m≦W)": (58, Decimal("1391.4"), 19),
    "路肩盛土": (61, Decimal("98.2"), 16),
    "路体外盛土(W＜2.5m)": (11, Decimal("6.7"), 66),
    "路体外盛土(2.5m≦W＜4.0m)": (1, Decimal("1.9"), 76),
    "路体外盛土(4.0m≦W)": (11, Decimal("38.0"), 66),
    "畦畔盛土": (23, Decimal("8.6"), 54),
    "切土法面整形(左)": (13, Decimal("33.4"), 64),
    "切土法面整形(右)": (13, Decimal("44.9"), 64),
    "盛土法面整形(左)": (47, Decimal("511.3"), 30),
    "盛土法面整形(右)": (47, Decimal("502.7"), 30),
}
MAINLINE_WORK_EXPECTED_STATS = {
    "床掘": (55, Decimal("125.7"), 22),
    "埋戻(C)": (11, Decimal("34.9"), 66),
    "埋戻(D)": (46, Decimal("45.5"), 31),
}
LEGACY_WORK_EXPECTED_STATS = {
    "床掘": (25, Decimal("76.7"), 0),
    "埋戻(C)": (0, Decimal("0"), 25),
    "埋戻(D)": (25, Decimal("44.9"), 0),
}
MAINLINE_EXPECTED_DIGEST = (
    "7297e6e9d4664d47c0437e72a907576"
    "b65f7db0e35d691d8e5d6479ff4a10d74"
)
CATEGORY_TABLE_EXPECTED_DIGEST = (
    "cd2f1871d60167fa4cf031203d0274b0"
    "1b2c5dc4c3fee05009eebcdc336e6ad2"
)


def text_token(text, x, y, layer, height=0.5):
    return {
        "text": text,
        "x": x,
        "y": y,
        "anchor_x": x,
        "anchor_y": y,
        "layer": layer,
        "height": height,
        "handle": None,
    }


class TableDetectionTests(unittest.TestCase):
    def test_detects_two_block_header_signature(self):
        texts = [
            text_token(word, index * 3.0, 100.0, "D-MTR-TXT")
            for index, word in enumerate(app.EARTHWORK_HEADER_SEQUENCE)
        ]

        tables = app.detect_earthwork_tables(texts)

        self.assertEqual(1, len(tables))
        self.assertEqual("D-MTR-TXT", tables[0]["layer"])

    def test_detects_category_quantity_header_without_symbol_column(self):
        x_positions = (
            0.0,
            4.0,
            7.94,
            13.08,
            16.91,
            20.00,
        )
        texts = [
            text_token(
                word,
                x,
                100.0 + (0.03 if index % 2 else 0.0),
                "D-MTR-TXT",
            )
            for index, (word, x) in enumerate(
                zip(app.CATEGORY_TABLE_HEADER_SEQUENCE, x_positions)
            )
        ]

        tables = app.detect_earthwork_tables(texts)

        self.assertEqual(1, len(tables))
        self.assertEqual(app.CATEGORY_TABLE_FORMAT, tables[0]["format"])
        self.assertEqual(app.CATEGORY_TABLE_LINE_LAYER, tables[0]["line_layer"])

    def test_center_marker_match_is_exact(self):
        texts = [
            text_token("CL", 0.0, 0.0, "D-BMK"),
            text_token("C.L.", 0.0, 1.0, "D-BMK"),
            # 土工表の記号値であり、中心線アンカーではない。
            text_token("CL1-1", 0.0, 2.0, "D-BMK"),
            text_token("CL1-2", 0.0, 3.0, "D-BMK"),
            text_token("CL", 0.0, 4.0, "D-MTR-TXT"),
        ]

        markers = app.extract_center_markers(texts)

        self.assertEqual(["CL", "C.L."], [marker["text"] for marker in markers])

    def test_selects_only_explicit_mainline_tables(self):
        tables = [
            {"id": route, "layer": "D-MTR-TXT"}
            for route in ("本線", "Bランプ", "Cランプ")
        ]
        regions = [
            {
                "main_bounds": {
                    "min_x": index * 100.0,
                    "max_x": index * 100.0 + 20.0,
                    "min_y": 0.0,
                    "max_y": 10.0,
                },
                "text_height": 1.0,
            }
            for index in range(3)
        ]
        texts = [
            text_token(
                route,
                index * 100.0 + 5.0,
                11.0,
                "D-MTR-TXT",
                height=1.0,
            )
            for index, route in enumerate(
                ("本線", "Bランプ", "Cランプ")
            )
        ]

        selected, selected_regions, alignments = (
            app._select_mainline_tables(
                tables,
                regions,
                texts,
            )
        )

        self.assertEqual(["本線"], [table["id"] for table in selected])
        self.assertEqual([regions[0]], selected_regions)
        self.assertEqual(["本線"], alignments)

    def test_treats_all_unlabelled_tables_as_mainline(self):
        tables = [
            {"id": index, "layer": "D-MTR-TXT"}
            for index in range(2)
        ]
        regions = [
            {
                "main_bounds": {
                    "min_x": index * 100.0,
                    "max_x": index * 100.0 + 20.0,
                    "min_y": 0.0,
                    "max_y": 10.0,
                },
                "text_height": 1.0,
            }
            for index in range(2)
        ]

        selected, selected_regions, alignments = (
            app._select_mainline_tables(
                tables,
                regions,
                [],
            )
        )

        self.assertEqual(tables, selected)
        self.assertEqual(regions, selected_regions)
        self.assertEqual(["本線", "本線"], alignments)

    def test_rejects_mixed_labelled_and_unlabelled_tables(self):
        tables = [
            {"id": index, "layer": "D-MTR-TXT"}
            for index in range(2)
        ]
        regions = [
            {
                "main_bounds": {
                    "min_x": index * 100.0,
                    "max_x": index * 100.0 + 20.0,
                    "min_y": 0.0,
                    "max_y": 10.0,
                },
                "text_height": 1.0,
            }
            for index in range(2)
        ]
        texts = [
            text_token(
                "本線",
                5.0,
                11.0,
                "D-MTR-TXT",
                height=1.0,
            )
        ]

        with self.assertRaisesRegex(app.ExtractionError, "混在"):
            app._select_mainline_tables(
                tables,
                regions,
                texts,
            )

    def test_resolves_all_route_tables_before_discarding_ramps(self):
        tables = [
            {"id": "mainline", "x": 0.0, "y": 0.0},
            {"id": "ramp", "x": 100.0, "y": 0.0},
        ]
        regions = [{"id": "mainline"}, {"id": "ramp"}]
        data_groups = [
            [{"x": 0.0, "y": 0.0, "種別": "mainline"}],
            [{"x": 100.0, "y": 0.0, "種別": "ramp"}],
        ]
        stations = [
            {
                "text": "1",
                "plus": None,
                "x": 0.0,
                "y": 0.0,
                "anchor_y": 0.0,
            },
            {
                "text": "2",
                "plus": None,
                "x": 100.0,
                "y": 0.0,
                "anchor_y": 0.0,
            },
        ]

        with (
            patch.object(
                app,
                "detect_earthwork_tables",
                return_value=tables,
            ),
            patch.object(
                app,
                "detect_table_regions",
                return_value=regions,
            ),
            patch.object(
                app,
                "_detect_table_route_names",
                return_value=["本線", "Bランプ"],
            ),
            patch.object(
                app,
                "build_table_data_groups",
                return_value=data_groups,
            ),
            patch.object(app, "extract_no_texts", return_value=stations),
            patch.object(app, "extract_center_markers", return_value=[]),
        ):
            record_sets = app.extract_output_record_sets([], [])

        self.assertEqual({None}, set(record_sets))
        self.assertEqual(1, len(record_sets[None]))
        self.assertEqual(Decimal("1"), record_sets[None][0]["測点"])
        self.assertEqual("mainline", record_sets[None][0]["データ"][0]["種別"])

    def test_rejects_data_group_far_from_its_only_table(self):
        tables = [{"x": 0.0, "y": 0.0}]
        data_groups = [[{"x": 100.0, "y": 0.0}]]

        with self.assertRaisesRegex(app.ExtractionError, "安全上限"):
            app.assign_data_groups_to_tables(tables, data_groups)

    def test_combines_upper_and_lower_roadbed_quantities(self):
        self.assertEqual(
            Decimal("25.2"),
            app._combine_table_quantities(
                Decimal("7.1"),
                Decimal("18.1"),
            ),
        )
        self.assertEqual("-", app._combine_table_quantities("-", "-"))

    def test_rejects_an_unexpected_hierarchical_source_schema(self):
        damaged_schema = list(app.HIERARCHICAL_SOURCE_SCHEMA)
        damaged_schema[9] = ("上部路床", "W≧4.0m")

        with self.assertRaisesRegex(app.ExtractionError, "種別・区分構成"):
            app._validate_hierarchical_source_schema(damaged_schema)


class StationResolutionTests(unittest.TestCase):
    def test_infers_20_or_100_meter_station_interval(self):
        standard = [
            {"測点": Decimal("283"), "追加距離": None},
            {"測点": Decimal("284"), "追加距離": None},
        ]
        hundred_meter = [
            {"測点": Decimal("7"), "追加距離": Decimal("80")},
            {"測点": Decimal("8"), "追加距離": None},
        ]

        self.assertEqual(
            Decimal("20"),
            app.infer_station_interval(standard),
        )
        self.assertEqual(
            Decimal("100"),
            app.infer_station_interval(hundred_meter),
        )

    def test_uses_single_station_assigned_to_each_table_when_cl_is_absent(self):
        tables = [
            {"x": 0.0, "y": 0.0},
            {"x": 100.0, "y": 0.0},
        ]
        stations = [
            {"text": "1", "plus": None, "x": 10.0, "y": 0.0},
            {"text": "2", "plus": None, "x": 90.0, "y": 0.0},
        ]

        resolved = app.resolve_table_stations(tables, stations, [])

        self.assertEqual(["1", "2"], [station["text"] for station in resolved])

    def test_uses_cl_before_direct_table_to_station_distance(self):
        tables = [{"x": 60.0, "y": 0.0}]
        center_markers = [{"text": "CL", "x": 0.0, "y": 0.0}]
        stations = [
            {"text": "10", "plus": Decimal("20"), "x": 0.0, "y": 5.0},
            {"text": "99", "plus": None, "x": 55.0, "y": 0.0},
        ]

        resolved = app.resolve_table_stations(tables, stations, center_markers)

        self.assertEqual("10", resolved[0]["text"])

    def test_rejects_multiple_local_stations_without_cl(self):
        tables = [{"x": 0.0, "y": 0.0}]
        stations = [
            {"text": "1", "plus": None, "x": 1.0, "y": 0.0},
            {"text": "2", "plus": None, "x": 2.0, "y": 0.0},
        ]

        with self.assertRaises(app.ExtractionError):
            app.resolve_table_stations(tables, stations, [])

    def test_rejects_nearly_equal_cl_station_candidates(self):
        tables = [{"x": 60.0, "y": 0.0}]
        center_markers = [{"text": "CL", "x": 0.0, "y": 0.0}]
        stations = [
            {"text": "1", "plus": None, "x": -5.0, "y": 0.0},
            {"text": "2", "plus": None, "x": 5.0, "y": 0.0},
        ]

        with self.assertRaisesRegex(app.ExtractionError, "測点候補"):
            app.resolve_table_stations(tables, stations, center_markers)

    def test_rejects_single_far_station_without_cl(self):
        tables = [{"x": 0.0, "y": 0.0}]
        stations = [
            {"text": "1", "plus": None, "x": 1000.0, "y": 0.0},
        ]

        with self.assertRaisesRegex(app.ExtractionError, "安全上限"):
            app.resolve_table_stations(tables, stations, [])


class VerifiedLocalRegressionTests(unittest.TestCase):
    good_dxf = LOCAL_CASE_DIR / "05横断図283～306.dxf"
    verified_workbook = (
        VERIFIED_OUTPUT_DIR / "01-01-03土工_05横断図283～306.xlsx"
    )
    problem_dxf = LOCAL_CASE_DIR / "05_本線横断図.dxf"
    category_table_dxf = PROJECT_ROOT / "input" / "004_横断図240514.dxf"

    def _first_good_table_texts(self):
        if not self.good_dxf.exists():
            self.skipTest("ローカルの検証用DXFがありません。")

        doc = ezdxf.readfile(self.good_dxf)
        modelspace = doc.modelspace()
        texts = app.collect_all_texts(modelspace)
        lines = app.collect_all_lines(modelspace)
        tables = app.detect_earthwork_tables(texts)
        regions = app.detect_table_regions(tables, lines, texts)
        return app._texts_in_table_region(
            texts,
            tables[0],
            regions[0],
        )

    def test_legacy_table_rejects_a_missing_unit_row(self):
        table_texts = self._first_good_table_texts()
        unit = next(
            text for text in table_texts
            if app.UNIT_PATTERN.fullmatch(text["text"])
        )
        damaged_texts = [
            text for text in table_texts
            if text is not unit
        ]

        with self.assertRaises(app.ExtractionError):
            app._build_legacy_table_group(damaged_texts)

    def test_legacy_table_rejects_a_missing_quantity(self):
        table_texts = self._first_good_table_texts()
        quantities = [
            text for text in table_texts
            if text["text"] == "9.2"
        ]
        self.assertEqual(1, len(quantities))
        damaged_texts = [
            text for text in table_texts
            if text is not quantities[0]
        ]

        with self.assertRaisesRegex(app.ExtractionError, "数値化"):
            app._build_legacy_table_group(damaged_texts)

    def test_input_layout_migration_is_idempotent(self):
        if not self.verified_workbook.exists():
            self.skipTest("旧配置の確認済みExcelがありません。")

        workbook = load_workbook(self.verified_workbook, data_only=False)
        try:
            sheet = workbook[app.SHEET_NAME]
            self.assertEqual(34, sheet.max_column)

            app._prepare_input_sheet_layout(sheet)
            self.assertEqual(35, sheet.max_column)
            self.assertEqual(
                list(app.WORK_OUTPUT_KINDS),
                [sheet[f"{column}2"].value for column in ("AC", "AD", "AE")],
            )

            app._prepare_input_sheet_layout(sheet)
            self.assertEqual(35, sheet.max_column)

            with tempfile.TemporaryDirectory() as temporary_directory:
                migrated = Path(temporary_directory) / "migrated.xlsx"
                workbook.save(migrated)
                reloaded = load_workbook(migrated, data_only=False)
                try:
                    reloaded_sheet = reloaded[app.SHEET_NAME]
                    self.assertEqual(35, reloaded_sheet.max_column)
                    for column in ("AC", "AD", "AE"):
                        self.assertFalse(
                            reloaded_sheet.column_dimensions[column].hidden
                        )
                        self.assertEqual(
                            reloaded_sheet.column_dimensions["AB"].width,
                            reloaded_sheet.column_dimensions[column].width,
                        )
                        for row in range(2, 6):
                            self.assertEqual(
                                reloaded_sheet[f"AB{row}"].style_id,
                                reloaded_sheet[f"{column}{row}"].style_id,
                            )
                    app._prepare_input_sheet_layout(reloaded_sheet)
                    self.assertEqual(35, reloaded_sheet.max_column)
                finally:
                    reloaded.close()
        finally:
            workbook.close()

    def test_short_output_clears_unused_template_detail_rows(self):
        template = PROJECT_ROOT / "template" / app.TEMPLATE_NAME
        if not template.exists():
            self.skipTest("Excelテンプレートがありません。")

        records = [
            {
                "測点": Decimal("1"),
                "追加距離": None,
                "データ": [],
            }
        ]
        with tempfile.TemporaryDirectory() as temporary_directory:
            output = Path(temporary_directory) / "short.xlsx"
            app.write_output_workbook(template, output, records)
            workbook = load_workbook(output, data_only=False)
            try:
                work_sheet = workbook["作業土工1"]
                self.assertTrue(
                    all(
                        cell.value is None
                        for row in range(7, 41)
                        for cell in work_sheet[row]
                    )
                )
                input_sheet = workbook[app.SHEET_NAME]
                self.assertEqual(35, input_sheet.max_column)
                transfer_columns = (
                    3,
                    5,
                    7,
                    11,
                    14,
                    15,
                    17,
                    18,
                    19,
                    20,
                    21,
                    22,
                    23,
                    24,
                    25,
                    26,
                    27,
                    28,
                    29,
                    30,
                    31,
                    32,
                    33,
                    34,
                    35,
                )
                self.assertTrue(
                    all(
                        input_sheet.cell(row, column).value is None
                        for row in range(6, input_sheet.max_row + 1)
                        for column in transfer_columns
                    )
                )
            finally:
                workbook.close()

    def test_verified_good_dxf_matches_workbook_cell_values(self):
        if not self.good_dxf.exists() or not self.verified_workbook.exists():
            self.skipTest("ローカルの検証用DXFまたは確認済みExcelがありません。")

        doc = ezdxf.readfile(self.good_dxf)
        modelspace = doc.modelspace()
        texts = app.collect_all_texts(modelspace)
        lines = app.collect_all_lines(modelspace)
        self.assertEqual(25, len(app.detect_earthwork_tables(texts)))
        self.assertEqual(25, len(app.extract_no_texts(texts)))
        self.assertEqual(0, len(app.extract_center_markers(texts)))

        tables = app.detect_earthwork_tables(texts)
        regions = app.detect_table_regions(tables, lines, texts)
        self.assertEqual({37}, {region["line_count"] for region in regions})
        self.assertEqual({1}, {region["attached_count"] for region in regions})

        records = app.extract_output_records(texts, lines)
        self.assertEqual(25, len(records))

        actual_work_stats = {}
        for kind in app.WORK_OUTPUT_KINDS:
            values = [
                data["数量"]
                for record in records
                for data in record["データ"]
                if data["種別"] == kind
            ]
            numeric = [
                value for value in values
                if isinstance(value, Decimal)
            ]
            actual_work_stats[kind] = (
                len(numeric),
                sum(numeric, Decimal("0")),
                values.count(None),
            )
        self.assertEqual(LEGACY_WORK_EXPECTED_STATS, actual_work_stats)

        sample = next(
            record for record in records
            if record["測点"] == Decimal("305")
            and record["追加距離"] is None
        )
        sample_values = {
            data["種別"]: data["数量"]
            for data in sample["データ"]
        }
        self.assertEqual(Decimal("6.9"), sample_values["床掘"])
        self.assertIsNone(sample_values["埋戻(C)"])
        self.assertEqual(Decimal("3.2"), sample_values["埋戻(D)"])

        with tempfile.TemporaryDirectory() as temporary_directory:
            generated_path = Path(temporary_directory) / "generated.xlsx"
            app.write_output_workbook(
                PROJECT_ROOT / "template" / app.TEMPLATE_NAME,
                generated_path,
                records,
            )
            expected = load_workbook(self.verified_workbook, data_only=False)
            actual = load_workbook(generated_path, data_only=False)
            try:
                expected_sheet = expected[app.SHEET_NAME]
                actual_sheet = actual[app.SHEET_NAME]
                mismatches = [
                    (
                        row,
                        column,
                        expected_sheet.cell(row, column).value,
                        actual_sheet.cell(row, column).value,
                    )
                    for row in range(1, expected_sheet.max_row + 1)
                    for column in range(1, 29)
                    if expected_sheet.cell(row, column).value
                    != actual_sheet.cell(row, column).value
                ]

                self.assertEqual(
                    list(app.WORK_OUTPUT_KINDS),
                    [actual_sheet[f"{column}2"].value for column in ("AC", "AD", "AE")],
                )
                self.assertEqual(
                    ["断面積"] * 3,
                    [actual_sheet[f"{column}3"].value for column in ("AC", "AD", "AE")],
                )
                self.assertEqual(
                    ["(㎡)"] * 3,
                    [actual_sheet[f"{column}4"].value for column in ("AC", "AD", "AE")],
                )
                for row in range(
                    app.ROW_START,
                    app.ROW_START + len(records),
                ):
                    self.assertEqual(
                        f'=IF(ISBLANK(C{row}),"",G{row})',
                        actual_sheet[f"AF{row}"].value,
                    )
                    self.assertEqual(
                        f'=IF(ISBLANK(C{row}),"",SUM(K{row}:X{row}))',
                        actual_sheet[f"AG{row}"].value,
                    )
                    self.assertEqual(
                        f'=IF(ISBLANK(C{row}),"",SUM(Y{row}:Z{row}))',
                        actual_sheet[f"AH{row}"].value,
                    )
                    self.assertEqual(
                        f'=IF(ISBLANK(C{row}),"",SUM(AA{row}:AB{row}))',
                        actual_sheet[f"AI{row}"].value,
                    )
                self.assertTrue(
                    all(
                        actual_sheet[f"{column}{row}"].value is None
                        for row in range(
                            app.ROW_START + len(records),
                            actual_sheet.max_row + 1,
                        )
                        for column in app.INPUT_SUMMARY_COLUMNS.values()
                    )
                )

                for input_row, record in enumerate(records, start=app.ROW_START):
                    quantities = {
                        data["種別"]: data["数量"]
                        for data in record["データ"]
                    }
                    for kind, column in app.INPUT_WORK_COLUMNS.items():
                        expected_value = quantities[kind]
                        actual_value = actual_sheet[f"{column}{input_row}"].value
                        if isinstance(expected_value, Decimal):
                            actual_value = Decimal(str(actual_value))
                        self.assertEqual(expected_value, actual_value)

                work_sheet = actual["作業土工1"]
                for index in range(len(records)):
                    input_row = app.ROW_START + index
                    work_row = app.ROW_START + 1 + index
                    for output_column, input_column in (
                        ("G", "AC"),
                        ("J", "AE"),
                        ("M", "AD"),
                    ):
                        self.assertEqual(
                            f'=IF(入力!{input_column}{input_row}="","",'
                            f"入力!{input_column}{input_row})",
                            work_sheet[f"{output_column}{work_row}"].value,
                        )

                check_sheet = actual["チェック"]
                self.assertEqual("作業土工", check_sheet["B25"].value)
                self.assertEqual(
                    '=IF(ISBLANK(INDEX(入力!$AC$5:$AC$29,$C$2)),"",'
                    "INDEX(入力!$AC$5:$AC$29,$C$2))",
                    check_sheet["D25"].value,
                )
                self.assertEqual(
                    '=IF(ISBLANK(INDEX(入力!$AE$5:$AE$29,$C$2)),"",'
                    "INDEX(入力!$AE$5:$AE$29,$C$2))",
                    check_sheet["D27"].value,
                )

                ref_formulas = [
                    (sheet.title, cell.coordinate, cell.value)
                    for sheet in actual.worksheets
                    for row in sheet.iter_rows()
                    for cell in row
                    if (
                        isinstance(cell.value, str)
                        and cell.value.startswith("=")
                        and "#REF!" in cell.value
                    )
                ]
            finally:
                expected.close()
                actual.close()

        self.assertEqual([], mismatches)
        self.assertEqual([], ref_formulas)

    def test_mainline_dxf_produces_77_complete_records(self):
        if not self.problem_dxf.exists():
            self.skipTest("ローカルの検証用DXFがありません。")

        doc = ezdxf.readfile(self.problem_dxf)
        modelspace = doc.modelspace()
        texts = app.collect_all_texts(modelspace)
        lines = app.collect_all_lines(modelspace)
        tables = app.detect_earthwork_tables(texts)
        stations = app.extract_no_texts(texts)
        center_markers = app.extract_center_markers(texts)
        self.assertEqual(77, len(tables))
        self.assertEqual(148, len(stations))
        self.assertEqual(141, len(center_markers))

        resolved_stations = app.resolve_table_stations(
            tables,
            stations,
            center_markers,
        )
        self.assertEqual(77, len(resolved_stations))
        self.assertEqual(
            77,
            len({station["handle"] for station in resolved_stations}),
        )
        first_low_table_index = min(
            range(len(tables)),
            key=lambda index: abs(tables[index]["y"] - 1260.654),
        )
        self.assertEqual(
            ("7", Decimal("40")),
            (
                resolved_stations[first_low_table_index]["text"],
                resolved_stations[first_low_table_index]["plus"],
            ),
        )

        regions = app.detect_table_regions(tables, lines, texts)
        self.assertEqual(77, len(regions))
        self.assertEqual({45}, {region["line_count"] for region in regions})
        self.assertEqual({0}, {region["attached_count"] for region in regions})

        records = app.extract_output_records(texts, lines)
        self.assertEqual(77, len(records))
        self.assertEqual(
            77,
            len({(record["測点"], record["追加距離"]) for record in records}),
        )
        self.assertTrue(
            all(
                [data["種別"] for data in record["データ"]]
                == list(app.HIERARCHICAL_ALL_OUTPUT_KINDS)
                for record in records
            )
        )

        actual_stats = {}
        for kind in app.HIERARCHICAL_OUTPUT_KINDS:
            values = [
                data["数量"]
                for record in records
                for data in record["データ"]
                if data["種別"] == kind
            ]
            numeric = [
                value for value in values
                if isinstance(value, Decimal)
            ]
            actual_stats[kind] = (
                len(numeric),
                sum(numeric, Decimal("0")),
                values.count("-"),
            )
        self.assertEqual(BAD_EXPECTED_STATS, actual_stats)

        actual_work_stats = {}
        for kind in app.WORK_OUTPUT_KINDS:
            values = [
                data["数量"]
                for record in records
                for data in record["データ"]
                if data["種別"] == kind
            ]
            numeric = [
                value for value in values
                if isinstance(value, Decimal)
            ]
            actual_work_stats[kind] = (
                len(numeric),
                sum(numeric, Decimal("0")),
                values.count("-"),
            )
        self.assertEqual(
            MAINLINE_WORK_EXPECTED_STATS,
            actual_work_stats,
        )

        digest_payload = [
            [
                str(record["測点"]),
                (
                    ""
                    if record["追加距離"] is None
                    else str(record["追加距離"])
                ),
                [
                    [data["種別"], str(data["数量"])]
                    for data in record["データ"]
                ],
            ]
            for record in records
        ]
        digest = hashlib.sha256(
            json.dumps(
                digest_payload,
                ensure_ascii=False,
                separators=(",", ":"),
            ).encode()
        ).hexdigest()
        self.assertEqual(MAINLINE_EXPECTED_DIGEST, digest)

        sample = next(
            record for record in records
            if record["測点"] == Decimal("7")
            and record["追加距離"] == Decimal("40")
        )
        sample_values = {
            data["種別"]: data["数量"]
            for data in sample["データ"]
        }
        self.assertEqual(Decimal("236.3"), sample_values["オープン掘削"])
        self.assertEqual(Decimal("0.4"), sample_values["路肩盛土"])
        self.assertEqual(
            Decimal("10.7"),
            sample_values["切土法面整形(左)"],
        )
        self.assertEqual(
            Decimal("11.2"),
            sample_values["切土法面整形(右)"],
        )
        self.assertEqual(Decimal("2.7"), sample_values["床掘"])
        self.assertEqual("-", sample_values["埋戻(C)"])
        self.assertEqual(Decimal("2.1"), sample_values["埋戻(D)"])

        with tempfile.TemporaryDirectory() as temporary_directory:
            generated_path = (
                Path(temporary_directory)
                / "01-01-03土工_05_本線横断図.xlsx"
            )
            app.write_output_workbook(
                PROJECT_ROOT / "template" / app.TEMPLATE_NAME,
                generated_path,
                records,
            )
            workbook = load_workbook(generated_path, data_only=False)
            try:
                input_sheet = workbook[app.SHEET_NAME]
                self.assertEqual(100, input_sheet["E4"].value)
                self.assertEqual(
                    '=IF(ISBLANK(C81),"",SUM(AA81:AB81))',
                    input_sheet["AI81"].value,
                )

                self.assertEqual(
                    list(app.WORK_OUTPUT_KINDS),
                    [input_sheet[f"{column}2"].value for column in ("AC", "AD", "AE")],
                )
                for input_row, record in enumerate(records, start=app.ROW_START):
                    quantities = {
                        data["種別"]: data["数量"]
                        for data in record["データ"]
                    }
                    for kind, column in app.INPUT_WORK_COLUMNS.items():
                        expected = quantities[kind]
                        actual = input_sheet[f"{column}{input_row}"].value
                        if isinstance(expected, Decimal):
                            actual = Decimal(str(actual))
                        self.assertEqual(expected, actual)

                positions = [
                    (
                        record["測点"] * Decimal("100")
                        + (record["追加距離"] or Decimal("0"))
                    )
                    for record in records
                ]
                distances = [
                    current - previous
                    for previous, current in zip(
                        positions,
                        positions[1:],
                    )
                ]
                self.assertTrue(all(distance > 0 for distance in distances))
                self.assertEqual(Decimal("2400"), sum(distances))

                for sheet_name in (
                    "掘削",
                    "路体",
                    "路床",
                    "路体外",
                    "作業土工1",
                ):
                    sheet = workbook[sheet_name]
                    self.assertEqual("=入力!C81", sheet["C82"].value)
                    self.assertEqual(84, sheet.max_row)
                    self.assertIn("$84", sheet.print_area)

                slope_sheet = workbook["法面整形"]
                self.assertEqual("=入力!C81", slope_sheet["C82"].value)
                self.assertEqual("=I84+L84", slope_sheet["L85"].value)
                self.assertIn("$85", slope_sheet.print_area)

                roadbed_sheet = workbook["路床"]
                self.assertEqual("=入力!T81", roadbed_sheet["P82"].value)
                self.assertEqual("=R83", roadbed_sheet["R84"].value)

                work_sheet = workbook["作業土工1"]
                for index in range(len(records)):
                    input_row = app.ROW_START + index
                    work_row = app.ROW_START + 1 + index
                    for output_column, input_column in (
                        ("G", "AC"),
                        ("J", "AE"),
                        ("M", "AD"),
                    ):
                        self.assertEqual(
                            f'=IF(入力!{input_column}{input_row}="","",'
                            f"入力!{input_column}{input_row})",
                            work_sheet[f"{output_column}{work_row}"].value,
                        )

                self.assertEqual(
                    "=路床!R84",
                    workbook["集計表"]["F17"].value,
                )
                self.assertEqual(
                    "=作業土工1!O84",
                    workbook["作業土工2"]["H5"].value,
                )

                ref_formulas = [
                    (sheet.title, cell.coordinate, cell.value)
                    for sheet in workbook.worksheets
                    for row in sheet.iter_rows()
                    for cell in row
                    if (
                        isinstance(cell.value, str)
                        and cell.value.startswith("=")
                        and "#REF!" in cell.value
                    )
                ]
                self.assertEqual([], ref_formulas)

                check_sheet = workbook["チェック"]
                self.assertEqual(
                    '=IF(ISBLANK(INDEX(入力!$AB$5:$AB$81,$C$2)),"",'
                    "INDEX(入力!$AB$5:$AB$81,$C$2))",
                    check_sheet["D24"].value,
                )
                self.assertEqual(
                    '=IF(ISBLANK(INDEX(入力!$AC$5:$AC$81,$C$2)),"",'
                    "INDEX(入力!$AC$5:$AC$81,$C$2))",
                    check_sheet["D25"].value,
                )
                self.assertEqual(
                    '=IF(ISBLANK(INDEX(入力!$AE$5:$AE$81,$C$2)),"",'
                    "INDEX(入力!$AE$5:$AE$81,$C$2))",
                    check_sheet["D27"].value,
                )
                self.assertEqual(
                    "掘削!$C$6:$C$82",
                    workbook["掘削"].defined_names["_1B"].attr_text,
                )
            finally:
                workbook.close()

    def test_category_quantity_dxf_ignores_symbols_and_outputs_only_mainline(self):
        if not self.category_table_dxf.exists():
            self.skipTest("ローカルの種別型検証用DXFがありません。")

        doc = ezdxf.readfile(self.category_table_dxf)
        modelspace = doc.modelspace()
        texts = app.collect_all_texts(modelspace)
        lines = app.collect_all_lines(modelspace)
        tables = app.detect_earthwork_tables(texts)
        self.assertEqual(103, len(tables))
        self.assertEqual(
            {app.CATEGORY_TABLE_FORMAT},
            {table["format"] for table in tables},
        )

        regions = app.detect_table_regions(tables, lines, texts)
        self.assertEqual({27}, {region["line_count"] for region in regions})
        self.assertEqual({0}, {region["attached_count"] for region in regions})

        record_sets = app.extract_output_record_sets(texts, lines)
        self.assertEqual(
            {
                "本線": 52,
            },
            {
                alignment: len(records)
                for alignment, records in record_sets.items()
            },
        )
        self.assertEqual(
            {
                "本線": (Decimal("214"), Decimal("265")),
            },
            {
                alignment: (
                    records[0]["測点"],
                    records[-1]["測点"],
                )
                for alignment, records in record_sets.items()
            },
        )

        symbol_pattern = re.compile(
            r"^(?:CA\d+|BA\d+(?:-\d+)?|CL\d+|BL\d+|"
            r"W\d+(?:-\d+)?|L|A)$"
        )
        texts_with_scrambled_symbols = []
        for index, text in enumerate(texts):
            copied = dict(text)
            if (
                text["text"] == "記号"
                or symbol_pattern.fullmatch(text["text"])
            ):
                copied["text"] = f"不定記号{index}"
            texts_with_scrambled_symbols.append(copied)
        self.assertEqual(
            record_sets,
            app.extract_output_record_sets(
                texts_with_scrambled_symbols,
                lines,
            ),
        )

        for alignment, records in record_sets.items():
            self.assertTrue(
                all(record["路線"] == alignment for record in records)
            )
            self.assertTrue(
                all(
                    [data["種別"] for data in record["データ"]]
                    == list(app.HIERARCHICAL_OUTPUT_KINDS)
                    for record in records
                )
            )
            for record in records:
                quantities = {
                    data["種別"]: data["数量"]
                    for data in record["データ"]
                }
                self.assertIsNone(quantities["切土法面整形(右)"])
                self.assertIsNone(quantities["盛土法面整形(右)"])

        digest_payload = {
            alignment: [
                [
                    str(record["測点"]),
                    (
                        ""
                        if record["追加距離"] is None
                        else str(record["追加距離"])
                    ),
                    [
                        [
                            data["種別"],
                            (
                                ""
                                if data["数量"] is None
                                else str(data["数量"])
                            ),
                        ]
                        for data in record["データ"]
                    ],
                ]
                for record in records
            ]
            for alignment, records in record_sets.items()
        }
        digest = hashlib.sha256(
            json.dumps(
                digest_payload,
                ensure_ascii=False,
                separators=(",", ":"),
            ).encode()
        ).hexdigest()
        self.assertEqual(CATEGORY_TABLE_EXPECTED_DIGEST, digest)

        with tempfile.TemporaryDirectory() as temporary_directory:
            temporary_root = Path(temporary_directory)
            input_directory = temporary_root / "input"
            template_directory = temporary_root / "template"
            input_directory.mkdir()
            template_directory.mkdir()
            shutil.copy2(
                self.category_table_dxf,
                input_directory / self.category_table_dxf.name,
            )
            shutil.copy2(
                PROJECT_ROOT / "template" / app.TEMPLATE_NAME,
                template_directory / app.TEMPLATE_NAME,
            )

            original_get_base_path = app.get_base_path
            app.get_base_path = lambda: str(temporary_root)
            try:
                processed, failures = app.proc([])
            finally:
                app.get_base_path = original_get_base_path

            self.assertEqual([], failures)
            self.assertEqual([self.category_table_dxf.name], processed)
            output_directories = list(
                (temporary_root / "output").glob("exec_*")
            )
            self.assertEqual(1, len(output_directories))
            output_paths = {
                path.stem.rsplit("_", 1)[-1]: path
                for path in output_directories[0].glob("*.xlsx")
            }
            self.assertEqual(set(record_sets), set(output_paths))

            for alignment, records in record_sets.items():
                workbook = load_workbook(
                    output_paths[alignment],
                    data_only=False,
                )
                try:
                    input_sheet = workbook[app.SHEET_NAME]
                    self.assertEqual(20, input_sheet["E4"].value)
                    for index, record in enumerate(records, start=app.ROW_START):
                        quantities = {
                            data["種別"]: data["数量"]
                            for data in record["データ"]
                        }
                        cut_left = input_sheet.cell(index, 25).value
                        fill_left = input_sheet.cell(index, 27).value
                        if isinstance(
                            quantities["切土法面整形(左)"],
                            Decimal,
                        ):
                            cut_left = Decimal(str(cut_left))
                        if isinstance(
                            quantities["盛土法面整形(左)"],
                            Decimal,
                        ):
                            fill_left = Decimal(str(fill_left))
                        self.assertEqual(
                            quantities["切土法面整形(左)"],
                            cut_left,
                        )
                        self.assertIsNone(input_sheet.cell(index, 26).value)
                        self.assertEqual(
                            quantities["盛土法面整形(左)"],
                            fill_left,
                        )
                        self.assertIsNone(input_sheet.cell(index, 28).value)
                        self.assertTrue(
                            all(
                                input_sheet[f"{column}{index}"].value is None
                                for column in ("AC", "AD", "AE")
                            )
                        )

                    ref_formulas = [
                        cell.value
                        for sheet in workbook.worksheets
                        for row in sheet.iter_rows()
                        for cell in row
                        if (
                            isinstance(cell.value, str)
                            and cell.value.startswith("=")
                            and "#REF!" in cell.value
                        )
                    ]
                    self.assertEqual([], ref_formulas)

                    if alignment == "本線":
                        work_sheet = workbook["作業土工1"]
                        for index in range(len(records)):
                            input_row = app.ROW_START + index
                            work_row = app.ROW_START + 1 + index
                            for output_column, input_column in (
                                ("G", "AC"),
                                ("J", "AE"),
                                ("M", "AD"),
                            ):
                                self.assertEqual(
                                    f'=IF(入力!{input_column}{input_row}="","",'
                                    f"入力!{input_column}{input_row})",
                                    work_sheet[f"{output_column}{work_row}"].value,
                                )
                finally:
                    workbook.close()

            self.assertEqual(
                [self.category_table_dxf.name],
                [
                    path.name
                    for path in input_directory.glob("used_*/*.dxf")
                ],
            )

    def test_proc_creates_both_workbooks_and_archives_both_dxf_files(self):
        if (
            not self.good_dxf.exists()
            or not self.problem_dxf.exists()
            or not (PROJECT_ROOT / "template" / app.TEMPLATE_NAME).exists()
        ):
            self.skipTest("ローカルの検証用DXFまたはテンプレートがありません。")

        with tempfile.TemporaryDirectory() as temporary_directory:
            temporary_root = Path(temporary_directory)
            input_directory = temporary_root / "input"
            template_directory = temporary_root / "template"
            input_directory.mkdir()
            template_directory.mkdir()
            shutil.copy2(self.good_dxf, input_directory / self.good_dxf.name)
            shutil.copy2(
                self.problem_dxf,
                input_directory / self.problem_dxf.name,
            )
            shutil.copy2(
                PROJECT_ROOT / "template" / app.TEMPLATE_NAME,
                template_directory / app.TEMPLATE_NAME,
            )

            original_get_base_path = app.get_base_path
            app.get_base_path = lambda: str(temporary_root)
            try:
                processed, failures = app.proc([])
            finally:
                app.get_base_path = original_get_base_path

            self.assertEqual([], failures)
            self.assertEqual(
                {self.good_dxf.name, self.problem_dxf.name},
                set(processed),
            )

            output_directories = list(
                (temporary_root / "output").glob("exec_*")
            )
            self.assertEqual(1, len(output_directories))
            output_names = {
                path.name
                for path in output_directories[0].glob("*.xlsx")
            }
            self.assertEqual(
                {
                    "01-01-03土工_05横断図283～306.xlsx",
                    "01-01-03土工_05_本線横断図.xlsx",
                },
                output_names,
            )
            self.assertEqual([], list(input_directory.glob("*.dxf")))
            archived_names = {
                path.name
                for path in input_directory.glob("used_*/*.dxf")
            }
            self.assertEqual(
                {self.good_dxf.name, self.problem_dxf.name},
                archived_names,
            )

    def test_proc_keeps_processing_errors_in_the_input_folder(self):
        with tempfile.TemporaryDirectory() as temporary_directory:
            temporary_root = Path(temporary_directory)
            input_directory = temporary_root / "input"
            template_directory = temporary_root / "template"
            input_directory.mkdir()
            template_directory.mkdir()
            broken_dxf = input_directory / "broken.dxf"
            broken_dxf.write_text("DXFではありません", encoding="utf-8")
            shutil.copy2(
                PROJECT_ROOT / "template" / app.TEMPLATE_NAME,
                template_directory / app.TEMPLATE_NAME,
            )

            original_get_base_path = app.get_base_path
            app.get_base_path = lambda: str(temporary_root)
            try:
                processed, failures = app.proc([])
            finally:
                app.get_base_path = original_get_base_path

            self.assertEqual([], processed)
            self.assertEqual(1, len(failures))
            self.assertIn("broken.dxf", failures[0])
            self.assertTrue(broken_dxf.exists())


if __name__ == "__main__":
    unittest.main()
