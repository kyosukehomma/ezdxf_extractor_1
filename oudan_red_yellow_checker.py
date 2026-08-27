import os
import sys
import re
import glob
import shutil
import unicodedata
from collections import Counter
from copy import copy
from decimal import Decimal, InvalidOperation
from datetime import datetime
from math import hypot
import tkinter as tk
from tkinter import messagebox

import ezdxf
from openpyxl import load_workbook
from openpyxl.formula.translate import Translator
from openpyxl.styles import Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.pagebreak import Break

# ==================================================
# 定数定義
# ==================================================

UNIT_PATTERN = re.compile(r"^(m2|㎡|m)$")
STATION_PATTERN = re.compile(
    r"^(?:No\.|NO\.)(\d+)(?:\+(\d+(?:\.\d+)?))?$",
    re.IGNORECASE,
)
PARENTHESIZED_STATION_PATTERN = re.compile(
    r"\((?:No\.|NO\.)(\d+)(?:\+(\d+(?:\.\d+)?))?\)",
    re.IGNORECASE,
)

TOLERANCE_Y_ROW   =  0.5
TOLERANCE_X_GROUP =  0.5
TOLERANCE_Y_GROUP = 10.0
MERGE_Y_TOLERANCE =  1.0

EARTHWORK_TEXT_LAYER_PREFIX = "D-MTR-TXT"
EARTHWORK_HEADER_SEQUENCE = ("種別", "単位", "数量", "種別", "単位", "数量")
CATEGORY_TABLE_HEADER_SEQUENCE = (
    "種別",
    "区分",
    "数量",
    "種別",
    "区分",
    "数量",
)
CATEGORY_TABLE_LINE_LAYER = "D-MTR-LINE"
CATEGORY_TABLE_FORMAT = "category_quantity"
CATEGORY_TABLE_HEADER_X_SPAN_MIN = 19.5
CATEGORY_TABLE_HEADER_X_SPAN_MAX = 22.0
CATEGORY_TABLE_KIND_PAIR_X_MIN = 12.5
CATEGORY_TABLE_KIND_PAIR_X_MAX = 13.7
CATEGORY_TABLE_KIND_PAIR_Y_TOLERANCE = 0.2
CATEGORY_TABLE_HEADER_Y_TOLERANCE = 0.3
TABLE_HEADER_Y_TOLERANCE = 0.15
TABLE_HEADER_X_GAP = 50.0
MAX_DATA_GROUP_TO_TABLE_DISTANCE = 10.0
LINE_AXIS_TOLERANCE = 1e-7
LINE_CONNECTION_TOLERANCE_FACTOR = 0.05
MAX_LINE_TO_TABLE_DISTANCE_FACTOR = 50.0
TABLE_MIN_WIDTH_FACTOR = 40.0
TABLE_MAX_WIDTH_FACTOR = 70.0
TABLE_MIN_HEIGHT_FACTOR = 20.0
TABLE_MAX_HEIGHT_FACTOR = 40.0
ATTACHED_GRID_MAX_GAP_FACTOR = 4.0
ATTACHED_GRID_LEFT_TOLERANCE_FACTOR = 1.0
TABLE_REGION_TOLERANCE_FACTOR = 0.05

LEGACY_OUTPUT_KINDS = (
    "オープン掘削",
    "路体盛土(W＜2.5m)",
    "路体盛土(2.5m≦W＜4.0m)",
    "路体盛土(4.0m≦W)",
    "路床盛土(W＜2.5m)",
    "路床盛土(2.5m≦W＜4.0m)",
    "路床盛土(4.0m≦W)",
    "路肩盛土",
    "路体外盛土(W＜2.5m)",
    "路体外盛土(2.5m≦W＜4.0m)",
    "路体外盛土(4.0m≦W)",
    "畦畔盛土",
    "余盛(路体盛土)",
    "余盛(路床盛土)",
    "余盛(載荷盛土)",
    "床掘",
    "埋戻",
    "切土法面整形(左)",
    "切土法面整形(右)",
    "盛土法面整形(左)",
    "盛土法面整形(右)",
    "表層",
    "基層",
    "上層路盤",
    "下層路盤",
    "凍上抑制層",
    "路肩表層(左)",
    "路肩表層(右)",
    "路肩下層路盤(左)",
    "路肩下層路盤(右)",
    "サンドマット",
    "サンドマット(端部)",
)

HIERARCHICAL_WIDTH_SUFFIX = {
    "W＜2.5m": "(W＜2.5m)",
    "2.5m≦W＜4.0m": "(2.5m≦W＜4.0m)",
    "W≧4.0m": "(4.0m≦W)",
}
HIERARCHICAL_WIDTH_CATEGORY = {
    "路体盛土": "路体盛土",
    "上部路床": "路床盛土",
    "下部路床": "路床盛土",
    "路体外盛土": "路体外盛土",
}
HIERARCHICAL_SOURCE_SCHEMA = (
    ("オープン掘削", "土砂"),
    ("路体外盛土", "W≧4.0m"),
    ("路体盛土", "W≧4.0m"),
    ("路体外盛土", "2.5m≦W＜4.0m"),
    ("路体盛土", "2.5m≦W＜4.0m"),
    ("路体外盛土", "W＜2.5m"),
    ("路体盛土", "W＜2.5m"),
    ("床堀り", "床堀り"),
    ("上部路床", "W≧4.0m"),
    ("埋戻し", "C"),
    ("上部路床", "2.5m≦W＜4.0m"),
    ("埋戻し", "D"),
    ("上部路床", "W＜2.5m"),
    ("路床安定処理", "路床安定処理"),
    ("下部路床", "W≧4.0m"),
    ("下部路床", "2.5m≦W＜4.0m"),
    ("法面整形工", "切土部(土砂)"),
    ("下部路床", "W＜2.5m"),
    ("法面整形工", "盛土部(土砂)"),
    ("路肩盛土", "路肩盛土"),
    ("法面工", "切土部(土砂)"),
    ("畦畔盛土", "畦畔盛土"),
    ("法面工", "盛土部(土砂)"),
)
HIERARCHICAL_OUTPUT_KINDS = (
    "オープン掘削",
    "路体盛土(W＜2.5m)",
    "路体盛土(2.5m≦W＜4.0m)",
    "路体盛土(4.0m≦W)",
    "路床盛土(W＜2.5m)",
    "路床盛土(2.5m≦W＜4.0m)",
    "路床盛土(4.0m≦W)",
    "路肩盛土",
    "路体外盛土(W＜2.5m)",
    "路体外盛土(2.5m≦W＜4.0m)",
    "路体外盛土(4.0m≦W)",
    "畦畔盛土",
    "切土法面整形(左)",
    "切土法面整形(右)",
    "盛土法面整形(左)",
    "盛土法面整形(右)",
)
HIERARCHICAL_WORK_OUTPUT_KINDS = (
    "床掘",
    "埋戻(C)",
    "埋戻(D)",
)
HIERARCHICAL_ALL_OUTPUT_KINDS = (
    HIERARCHICAL_OUTPUT_KINDS
    + HIERARCHICAL_WORK_OUTPUT_KINDS
)
CATEGORY_TABLE_SOURCE_SCHEMA = (
    ("オープン掘削", ((None, "オープン掘削"),)),
    (
        "路体盛土",
        (
            ("(W＜2.5m)", "路体盛土(W＜2.5m)"),
            ("(2.5m≦W＜4.0m)", "路体盛土(2.5m≦W＜4.0m)"),
            ("(4.0m≦W)", "路体盛土(4.0m≦W)"),
        ),
    ),
    (
        "路床盛土",
        (
            ("(W＜2.5m)", "路床盛土(W＜2.5m)"),
            ("(2.5m≦W＜4.0m)", "路床盛土(2.5m≦W＜4.0m)"),
            ("(4.0m≦W)", "路床盛土(4.0m≦W)"),
        ),
    ),
    ("路肩盛土", ((None, "路肩盛土"),)),
    (
        "路体外盛土",
        (
            ("(W＜2.5m)", "路体外盛土(W＜2.5m)"),
            ("(2.5m≦W＜4.0m)", "路体外盛土(2.5m≦W＜4.0m)"),
            ("(4.0m≦W)", "路体外盛土(4.0m≦W)"),
        ),
    ),
    ("畦畔盛土", ((None, "畦畔盛土"),)),
    ("切土法面整形", ((None, "切土法面整形(左)"),)),
    ("盛土法面整形", ((None, "盛土法面整形(左)"),)),
)
CATEGORY_TABLE_RIGHT_OUTPUT_KINDS = (
    "切土法面整形(右)",
    "盛土法面整形(右)",
)
CATEGORY_TABLE_MAINLINE_STATION_MIN = 100

STATION_TEXT_LAYER = "D-BMK-HTXT"
CENTER_MARKER_LAYER = "D-BMK"
CENTER_MARKER_TEXTS = {"CL", "C.L", "C.L."}
ROUTE_NAME_PATTERN = re.compile(r"^(?:本線|.+ランプ)$")
MAX_TABLE_TO_CENTER_DISTANCE = 75.0
MAX_TABLE_TO_CENTER_Y_DISTANCE = 15.0
MAX_CENTER_TO_STATION_DISTANCE = 15.0
MAX_TABLE_TO_STATION_DISTANCE = 70.0
MAX_TABLE_TO_STATION_Y_DISTANCE = 15.0
MIN_NEAREST_DISTANCE_GAP = 1.0

SHEET_NAME = "入力"
ROW_START = 5
VERIFIED_LEGACY_RECORD_COUNT = 25
STATION_INTERVAL_CANDIDATES = (Decimal("20"), Decimal("100"))
TEMPLATE_BASE = "01-01-03土工"
TEMPLATE_NAME = TEMPLATE_BASE + ".xlsx"


class ExtractionError(RuntimeError):
    """誤った転記を避けるために抽出処理を中断するエラー。"""

# ==================================================
# ユーティリティ
# ==================================================

def get_base_path():
    if getattr(sys, "frozen", False):
        return os.path.dirname(sys.executable)
    return os.path.dirname(os.path.abspath(__file__))

def normalize_text(s: str) -> str:
    return re.sub(r"\s+", "", unicodedata.normalize("NFKC", s))

def replace_angle_brackets(s: str) -> str:
    return s.replace("<", "＜").replace(">", "＞")

def sort_key(item):
    plus = item["追加距離"]
    return (item["測点"], Decimal("-1") if plus is None else plus)

# ==================================================
# DXF → TEXT 抽出
# ==================================================

def collect_all_texts(msp):
    texts = []
    for t in msp.query("TEXT"):
        s = replace_angle_brackets(normalize_text(t.dxf.text))
        alignment_point = t.get_placement()[1]
        texts.append({
            "text": s,
            "x": t.dxf.insert[0],
            "y": t.dxf.insert[1],
            "anchor_x": alignment_point.x,
            "anchor_y": alignment_point.y,
            "layer": t.dxf.layer,
            "height": t.dxf.height,
            "handle": t.dxf.handle,
        })
    return texts


def collect_all_lines(msp):
    lines = []
    for line in msp.query("LINE"):
        start = line.dxf.start
        end = line.dxf.end
        lines.append({
            "x": (start.x + end.x) / 2,
            "y": (start.y + end.y) / 2,
            "start_x": start.x,
            "start_y": start.y,
            "end_x": end.x,
            "end_y": end.y,
            "layer": line.dxf.layer,
            "handle": line.dxf.handle,
        })
    return lines


def extract_no_texts(all_texts):
    results = []
    for t in all_texts:
        if t.get("layer") != STATION_TEXT_LAYER:
            continue

        norm = unicodedata.normalize("NFKC", t["text"])
        match = STATION_PATTERN.fullmatch(norm)
        if match is None:
            match = PARENTHESIZED_STATION_PATTERN.search(norm)
        if match is None:
            continue

        try:
            plus = Decimal(match.group(2)) if match.group(2) is not None else None
        except InvalidOperation:
            continue

        results.append({
            "text": match.group(1),
            "plus": plus,
            "x": t["x"],
            "y": t["y"],
            "anchor_x": t["anchor_x"],
            "anchor_y": t["anchor_y"],
            "layer": t["layer"],
            "handle": t.get("handle"),
        })
    return results


def extract_center_markers(all_texts):
    return [
        t for t in all_texts
        if t.get("layer") == CENTER_MARKER_LAYER
        and t["text"].upper() in CENTER_MARKER_TEXTS
    ]

# ==================================================
# 行構築
# ==================================================

def build_ba_records(all_texts):
    records = []

    for item in all_texts:
        if not UNIT_PATTERN.fullmatch(item["text"]):
            continue

        same_row = [
            t for t in all_texts
            if abs(t["y"] - item["y"]) <= TOLERANCE_Y_ROW
        ]
        same_row.sort(key=lambda x: x["x"])

        idx = [t["x"] for t in same_row].index(item["x"])
        kind = same_row[idx - 1]["text"] if idx - 1 >= 0 else ""
        qty  = same_row[idx + 1]["text"] if idx + 1 < len(same_row) else ""

        try:
            qty = Decimal(qty)
        except Exception:
            pass

        records.append({
            "x": item["x"],
            "y": item["y"],
            "種別": kind,
            "単位": item["text"],
            "数量": qty
        })

    return records

# ==================================================
# グループ化
# ==================================================

def group_by_x(ba_records):
    groups = []
    for r in ba_records:
        for g in groups:
            if abs(g["x_ref"] - r["x"]) <= TOLERANCE_X_GROUP:
                g["items"].append(r)
                break
        else:
            groups.append({"x_ref": r["x"], "items": [r]})
    return groups

def group_by_y(x_groups):
    final = []
    for g in x_groups:
        items = sorted(g["items"], key=lambda i: i["y"], reverse=True)
        subgroups, buf, last_y = [], [], None
        for it in items:
            if last_y is None or abs(last_y - it["y"]) <= TOLERANCE_Y_GROUP:
                buf.append(it)
            else:
                subgroups.append(buf)
                buf = [it]
            last_y = it["y"]
        if buf:
            subgroups.append(buf)
        final.append(subgroups)
    return final

# ==================================================
# 種別補正
# ==================================================

WIDTH_MAP = {
    "路体盛土(W＜2.5m)": "路体盛土",
    "路床盛土(W＜2.5m)": "路床盛土",
    "路体外盛土(W＜2.5m)": "路体外盛土",
}

SIDE_RIGHT_MAP = {
    "切土法面整形(左)": "切土法面整形",
    "盛土法面整形(左)": "盛土法面整形",
    "路肩表層(左)": "路肩表層",
    "路肩下層路盤(左)": "路肩下層路盤",
}

def normalize_kinds(final_groups):
    for x_group in final_groups:
        for sg in x_group:
            if len(sg) < 2:
                continue

            base = sg[0]["種別"]

            # 幅員系
            if base == "オープン掘削":
                for i in range(1, len(sg)):
                    k = sg[i]["種別"]
                    if k in WIDTH_MAP:
                        prefix = WIDTH_MAP[k]
                        if i + 1 < len(sg) and sg[i + 1]["種別"].startswith("("):
                            sg[i + 1]["種別"] = f"{prefix}{sg[i + 1]['種別']}"
                        if i + 2 < len(sg) and sg[i + 2]["種別"].startswith("("):
                            sg[i + 2]["種別"] = f"{prefix}{sg[i + 2]['種別']}"

            # 左右系
            elif base == "切土法面整形(左)":
                for i in range(1, len(sg)):
                    prev, curr = sg[i - 1]["種別"], sg[i]["種別"]
                    if curr == "(右)" and prev in SIDE_RIGHT_MAP:
                        sg[i]["種別"] = f"{SIDE_RIGHT_MAP[prev]}(右)"

# ==================================================
# merge
# ==================================================

def merge_x_groups(final_groups):
    result, prev = [], None

    for g in final_groups:
        if prev is None:
            prev = g
            continue

        len_ok  = len(prev) == len(g)
        kind_ok = prev[0][0]["種別"] != g[0][0]["種別"]
        y_ok    = abs(prev[0][0]["y"] - g[0][0]["y"]) <= MERGE_Y_TOLERANCE

        if len_ok and kind_ok and y_ok:
            for i in range(len(prev)):
                prev[i].extend(g[i])
        else:
            result.append(prev)
            prev = g

    if prev:
        result.append(prev)
    return result


# ==================================================
# 土工表の検出・測点との対応付け
# ==================================================

def _distance(a, b):
    a_x = a.get("anchor_x", a["x"])
    a_y = a.get("anchor_y", a["y"])
    b_x = b.get("anchor_x", b["x"])
    b_y = b.get("anchor_y", b["y"])
    return hypot(a_x - b_x, a_y - b_y)


def _split_by_x_gap(items, max_gap):
    chunks = []
    current = []
    last_x = None

    for item in sorted(items, key=lambda t: t["x"]):
        if last_x is not None and item["x"] - last_x > max_gap:
            chunks.append(current)
            current = []
        current.append(item)
        last_x = item["x"]

    if current:
        chunks.append(current)
    return chunks


def _detect_unit_tables(all_texts):
    """左右2組の「種別・単位・数量」ヘッダーを検出する。"""
    header_words = set(EARTHWORK_HEADER_SEQUENCE)
    header_texts = [
        t for t in all_texts
        if t.get("layer", "").startswith(EARTHWORK_TEXT_LAYER_PREFIX)
        and t["text"] in header_words
    ]

    row_groups = []
    for item in sorted(
        header_texts,
        key=lambda t: (t["layer"], -t["y"], t["x"]),
    ):
        same_row = next(
            (
                group for group in row_groups
                if group["layer"] == item["layer"]
                and abs(group["y_ref"] - item["y"]) <= TABLE_HEADER_Y_TOLERANCE
            ),
            None,
        )
        if same_row is None:
            row_groups.append({
                "layer": item["layer"],
                "y_ref": item["y"],
                "items": [item],
            })
        else:
            same_row["items"].append(item)

    tables = []
    for row_group in row_groups:
        for chunk in _split_by_x_gap(
            row_group["items"],
            TABLE_HEADER_X_GAP,
        ):
            ordered_words = tuple(t["text"] for t in sorted(chunk, key=lambda t: t["x"]))
            if ordered_words != EARTHWORK_HEADER_SEQUENCE:
                continue

            tables.append({
                "x": sum(t["x"] for t in chunk) / len(chunk),
                "y": sum(t["y"] for t in chunk) / len(chunk),
                "layer": row_group["layer"],
                "line_layer": row_group["layer"],
                "format": "unit",
                "headers": tuple(chunk),
            })

    return tables


def _detect_category_quantity_tables(all_texts):
    """
    左右2組の「種別・区分・数量」ヘッダーを検出する。

    この形式は同じ表のヘッダー間でもY座標に僅かな差があるため、
    左右の「種別」の相対位置を起点に6項目を照合する。
    記号列のヘッダーと値は検出にも抽出にも使用しない。
    """
    header_words = set(CATEGORY_TABLE_HEADER_SEQUENCE)
    kind_headers = [
        text for text in all_texts
        if text.get("layer", "").startswith(EARTHWORK_TEXT_LAYER_PREFIX)
        and text["text"] == "種別"
    ]

    tables = []
    for left_kind in kind_headers:
        right_kinds = [
            text for text in kind_headers
            if text["layer"] == left_kind["layer"]
            and CATEGORY_TABLE_KIND_PAIR_X_MIN
            < text["x"] - left_kind["x"]
            < CATEGORY_TABLE_KIND_PAIR_X_MAX
            and abs(text["y"] - left_kind["y"])
            <= CATEGORY_TABLE_KIND_PAIR_Y_TOLERANCE
        ]
        if len(right_kinds) != 1:
            continue

        right_kind = right_kinds[0]
        candidates = sorted(
            [
                text for text in all_texts
                if text["layer"] == left_kind["layer"]
                and text["text"] in header_words
                and left_kind["x"] - CATEGORY_TABLE_KIND_PAIR_Y_TOLERANCE
                <= text["x"]
                <= right_kind["x"] + 7.5
                and abs(text["y"] - left_kind["y"])
                <= CATEGORY_TABLE_HEADER_Y_TOLERANCE
            ],
            key=lambda text: text["x"],
        )
        ordered_words = tuple(text["text"] for text in candidates)
        if ordered_words != CATEGORY_TABLE_HEADER_SEQUENCE:
            continue

        header_span = candidates[-1]["x"] - candidates[0]["x"]
        if not (
            CATEGORY_TABLE_HEADER_X_SPAN_MIN
            <= header_span
            <= CATEGORY_TABLE_HEADER_X_SPAN_MAX
        ):
            continue

        tables.append({
            "x": sum(text["x"] for text in candidates) / len(candidates),
            "y": sum(text["y"] for text in candidates) / len(candidates),
            "layer": left_kind["layer"],
            "line_layer": CATEGORY_TABLE_LINE_LAYER,
            "format": CATEGORY_TABLE_FORMAT,
            "headers": tuple(candidates),
        })

    return tables


def detect_earthwork_tables(all_texts):
    """対応する各ヘッダー形式から土工表を検出する。"""
    tables = (
        _detect_unit_tables(all_texts)
        + _detect_category_quantity_tables(all_texts)
    )
    tables.sort(key=lambda t: (-t["y"], t["x"]))
    return tables


def _line_bounds(line):
    return {
        "min_x": min(line["start_x"], line["end_x"]),
        "max_x": max(line["start_x"], line["end_x"]),
        "min_y": min(line["start_y"], line["end_y"]),
        "max_y": max(line["start_y"], line["end_y"]),
    }


def _lines_bounds(lines):
    return {
        "min_x": min(min(line["start_x"], line["end_x"]) for line in lines),
        "max_x": max(max(line["start_x"], line["end_x"]) for line in lines),
        "min_y": min(min(line["start_y"], line["end_y"]) for line in lines),
        "max_y": max(max(line["start_y"], line["end_y"]) for line in lines),
    }


def _point_in_bounds(x, y, bounds, tolerance=0.0):
    return (
        bounds["min_x"] - tolerance <= x <= bounds["max_x"] + tolerance
        and bounds["min_y"] - tolerance <= y <= bounds["max_y"] + tolerance
    )


def _line_segments_touch(first, second, tolerance):
    first_bounds = _line_bounds(first)
    second_bounds = _line_bounds(second)
    return not (
        first_bounds["max_x"] < second_bounds["min_x"] - tolerance
        or second_bounds["max_x"] < first_bounds["min_x"] - tolerance
        or first_bounds["max_y"] < second_bounds["min_y"] - tolerance
        or second_bounds["max_y"] < first_bounds["min_y"] - tolerance
    )


def _connected_line_components(lines, tolerance):
    components = []
    visited = set()

    for start_index in range(len(lines)):
        if start_index in visited:
            continue

        stack = [start_index]
        visited.add(start_index)
        component = []
        while stack:
            line_index = stack.pop()
            component.append(lines[line_index])
            for candidate_index in range(len(lines)):
                if candidate_index in visited:
                    continue
                if _line_segments_touch(
                    lines[line_index],
                    lines[candidate_index],
                    tolerance,
                ):
                    visited.add(candidate_index)
                    stack.append(candidate_index)

        components.append(component)

    return components


def detect_table_regions(tables, all_lines, all_texts):
    """形式ごとの罫線レイヤーから、表ごとの抽出領域を作る。"""
    if not tables:
        raise ExtractionError("抽出領域を作成する土工表がありません。")

    assignments = [[] for _ in tables]
    table_layers = {
        table.get("line_layer", table["layer"])
        for table in tables
    }
    component_line_layers = {
        table.get("line_layer", table["layer"])
        for table in tables
        if table.get("format") == CATEGORY_TABLE_FORMAT
    }

    for line in all_lines:
        if line["layer"] not in table_layers:
            continue
        if line["layer"] in component_line_layers:
            continue

        dx = abs(line["end_x"] - line["start_x"])
        dy = abs(line["end_y"] - line["start_y"])
        if dx > LINE_AXIS_TOLERANCE and dy > LINE_AXIS_TOLERANCE:
            raise ExtractionError("土工表レイヤーに斜めのLINEがあり、罫線を判定できません。")

        candidate_indexes = [
            index for index, table in enumerate(tables)
            if table.get("line_layer", table["layer"]) == line["layer"]
        ]
        table_index = min(
            candidate_indexes,
            key=lambda index: _distance(tables[index], line),
        )
        text_height = tables[table_index]["headers"][0]["height"]
        if (
            _distance(tables[table_index], line)
            > text_height * MAX_LINE_TO_TABLE_DISTANCE_FACTOR
        ):
            raise ExtractionError("土工表から離れた罫線が同じレイヤーに存在します。")
        assignments[table_index].append(line)

    for line_layer in component_line_layers:
        layer_tables = [
            (index, table)
            for index, table in enumerate(tables)
            if table.get("line_layer", table["layer"]) == line_layer
        ]
        layer_lines = [
            line for line in all_lines
            if line["layer"] == line_layer
        ]
        if not layer_lines:
            raise ExtractionError("種別型土工表の罫線を検出できませんでした。")

        for line in layer_lines:
            dx = abs(line["end_x"] - line["start_x"])
            dy = abs(line["end_y"] - line["start_y"])
            if dx > LINE_AXIS_TOLERANCE and dy > LINE_AXIS_TOLERANCE:
                raise ExtractionError(
                    "種別型土工表レイヤーに斜めのLINEがあります。"
                )

        text_heights = sorted(
            header["height"]
            for _, table in layer_tables
            for header in table["headers"]
        )
        text_height = text_heights[len(text_heights) // 2]
        connection_tolerance = (
            text_height * LINE_CONNECTION_TOLERANCE_FACTOR
        )
        components = _connected_line_components(
            layer_lines,
            connection_tolerance,
        )
        for component in components:
            bounds = _lines_bounds(component)
            matching_indexes = [
                index for index, table in layer_tables
                if all(
                    _point_in_bounds(
                        header["x"],
                        header["y"],
                        bounds,
                        connection_tolerance,
                    )
                    for header in table["headers"]
                )
            ]
            if len(matching_indexes) != 1:
                raise ExtractionError(
                    "種別型土工表の罫線をヘッダーへ一意に対応付け"
                    "できません。"
                )
            assignments[matching_indexes[0]].extend(component)

    regions = []
    for table_index, table in enumerate(tables):
        assigned_lines = assignments[table_index]
        if not assigned_lines:
            raise ExtractionError("土工表に対応する罫線を検出できませんでした。")

        heights = sorted(header["height"] for header in table["headers"])
        text_height = heights[len(heights) // 2]
        connection_tolerance = (
            text_height * LINE_CONNECTION_TOLERANCE_FACTOR
        )
        components = _connected_line_components(
            assigned_lines,
            connection_tolerance,
        )

        component_bounds = [
            _lines_bounds(component)
            for component in components
        ]
        main_indexes = [
            index for index, bounds in enumerate(component_bounds)
            if all(
                _point_in_bounds(
                    header["x"],
                    header["y"],
                    bounds,
                    connection_tolerance,
                )
                for header in table["headers"]
            )
        ]
        if len(main_indexes) != 1:
            raise ExtractionError(
                "土工表ヘッダーを含む主表罫線を一意に特定できません。"
            )

        main_index = main_indexes[0]
        main_bounds = component_bounds[main_index]
        main_width = main_bounds["max_x"] - main_bounds["min_x"]
        main_height = main_bounds["max_y"] - main_bounds["min_y"]
        if not (
            text_height * TABLE_MIN_WIDTH_FACTOR
            <= main_width
            <= text_height * TABLE_MAX_WIDTH_FACTOR
            and text_height * TABLE_MIN_HEIGHT_FACTOR
            <= main_height
            <= text_height * TABLE_MAX_HEIGHT_FACTOR
        ):
            raise ExtractionError("土工表罫線の幅または高さが想定範囲外です。")

        main_horizontal_count = sum(
            abs(line["end_y"] - line["start_y"]) <= LINE_AXIS_TOLERANCE
            for line in components[main_index]
        )
        main_vertical_count = sum(
            abs(line["end_x"] - line["start_x"]) <= LINE_AXIS_TOLERANCE
            for line in components[main_index]
        )
        if main_horizontal_count < 10 or main_vertical_count < 6:
            raise ExtractionError("土工表の水平・垂直罫線が不足しています。")

        selected_indexes = {main_index}
        attached_indexes = []
        for component_index, bounds in enumerate(component_bounds):
            if component_index == main_index:
                continue

            vertical_gap = main_bounds["min_y"] - bounds["max_y"]
            left_aligned = (
                abs(bounds["min_x"] - main_bounds["min_x"])
                <= text_height * ATTACHED_GRID_LEFT_TOLERANCE_FACTOR
            )
            within_width = (
                bounds["max_x"]
                <= main_bounds["max_x"] + text_height
            )
            captions = [
                text for text in all_texts
                if text["layer"] == table["layer"]
                and text["text"] == "作業土工"
                and bounds["min_x"] - text_height
                <= text["x"]
                <= bounds["max_x"] + text_height
                and abs(text["y"] - bounds["max_y"]) <= text_height * 2
            ]
            if (
                0 <= vertical_gap
                <= text_height * ATTACHED_GRID_MAX_GAP_FACTOR
                and left_aligned
                and within_width
                and len(captions) == 1
            ):
                horizontal_count = sum(
                    abs(line["end_y"] - line["start_y"])
                    <= LINE_AXIS_TOLERANCE
                    for line in components[component_index]
                )
                vertical_count = sum(
                    abs(line["end_x"] - line["start_x"])
                    <= LINE_AXIS_TOLERANCE
                    for line in components[component_index]
                )
                if horizontal_count >= 2 and vertical_count >= 2:
                    selected_indexes.add(component_index)
                    attached_indexes.append(component_index)

        if len(selected_indexes) != len(components):
            raise ExtractionError(
                "土工表付近に所属を判定できない罫線領域があります。"
            )

        regions.append({
            "areas": tuple(
                component_bounds[index]
                for index in sorted(selected_indexes)
            ),
            "main_bounds": main_bounds,
            "main_lines": tuple(components[main_index]),
            "line_count": sum(
                len(components[index])
                for index in selected_indexes
            ),
            "attached_count": len(attached_indexes),
            "text_height": text_height,
        })

    return regions


def _texts_in_table_region(all_texts, table, region):
    tolerance = (
        region["text_height"] * TABLE_REGION_TOLERANCE_FACTOR
    )
    return [
        text for text in all_texts
        if text["layer"] == table["layer"]
        and any(
            _point_in_bounds(
                text["x"],
                text["y"],
                bounds,
                tolerance,
            )
            for bounds in region["areas"]
        )
    ]


def _build_legacy_table_group(table_texts):
    ba_records = build_ba_records(table_texts)
    final = group_by_y(group_by_x(ba_records))
    normalize_kinds(final)
    final = merge_x_groups(final)
    groups = [
        subgroup for x_group in final for subgroup in x_group
    ]
    if len(groups) != 1:
        raise ExtractionError(
            f"1つの土工表から{len(groups)}個のデータ群が生成されました。"
        )
    group = groups[0]
    actual_kinds = tuple(record["種別"] for record in group)
    if actual_kinds != LEGACY_OUTPUT_KINDS:
        missing = [
            kind for kind in LEGACY_OUTPUT_KINDS
            if kind not in actual_kinds
        ]
        extra = [
            kind for kind in actual_kinds
            if kind not in LEGACY_OUTPUT_KINDS
        ]
        raise ExtractionError(
            f"既存型土工表の転記項目が一致しません"
            f"（件数={len(group)}, 不足={missing}, 余分={extra}）。"
        )

    invalid_quantities = [
        record["数量"] for record in group
        if not (
            isinstance(record["数量"], Decimal)
            or record["数量"] == "-"
        )
    ]
    if invalid_quantities:
        raise ExtractionError(
            f"既存型土工表に数値化できない数量があります"
            f"（{invalid_quantities}）。"
        )

    return group


def _unique_axis_values(values, tolerance):
    result = []
    for value in sorted(values):
        if not result or abs(result[-1] - value) > tolerance:
            result.append(value)
    return result


def _horizontal_cell_bounds(lines, x, y, tolerance):
    levels = _unique_axis_values(
        [
            (line["start_y"] + line["end_y"]) / 2
            for line in lines
            if abs(line["end_y"] - line["start_y"])
            <= LINE_AXIS_TOLERANCE
            and min(line["start_x"], line["end_x"]) - tolerance
            <= x
            <= max(line["start_x"], line["end_x"]) + tolerance
        ],
        tolerance,
    )
    lower = [level for level in levels if level <= y + tolerance]
    upper = [level for level in levels if level >= y - tolerance]
    if not lower or not upper:
        raise ExtractionError("表の結合セル境界を特定できませんでした。")
    return max(lower), min(upper)


def _parse_table_quantity(text):
    if text == "-":
        return "-"
    try:
        return Decimal(text)
    except InvalidOperation:
        return None


def _combine_table_quantities(first, second):
    values = [
        value for value in (first, second)
        if isinstance(value, Decimal)
    ]
    if not values:
        return "-"
    return sum(values, Decimal("0"))


def _hierarchical_category(
    record,
    table,
    region,
    table_texts,
    kind_headers,
    division_headers,
):
    main_bounds = region["main_bounds"]
    center_x = (main_bounds["min_x"] + main_bounds["max_x"]) / 2
    side_index = 0 if record["x"] < center_x else 1
    kind_x = kind_headers[side_index]["x"]
    division_x = division_headers[side_index]["x"]
    category_max_x = (kind_x + division_x) / 2
    side_min_x = (
        main_bounds["min_x"] if side_index == 0 else center_x
    )
    lower_y, upper_y = _horizontal_cell_bounds(
        region["main_lines"],
        kind_x,
        record["y"],
        region["text_height"] * LINE_CONNECTION_TOLERANCE_FACTOR,
    )
    candidates = [
        text for text in table_texts
        if side_min_x - region["text_height"] <= text["x"] < category_max_x
        and lower_y - region["text_height"] * 0.1
        <= text["y"]
        <= upper_y + region["text_height"] * 0.1
        and text["text"] != "種別"
    ]
    if len(candidates) != 1:
        raise ExtractionError(
            "階層型土工表の種別を一意に特定できませんでした。"
        )
    return candidates[0]["text"]


def _row_quantity_values(record, table_texts):
    values = []
    for text in sorted(table_texts, key=lambda item: item["x"]):
        if (
            text["x"] > record["x"]
            and abs(text["y"] - record["y"]) <= TOLERANCE_Y_ROW
        ):
            value = _parse_table_quantity(text["text"])
            if value is not None:
                values.append(value)
    return values


def _validate_hierarchical_source_schema(schema):
    if tuple(schema) != HIERARCHICAL_SOURCE_SCHEMA:
        raise ExtractionError(
            "階層型土工表の種別・区分構成が想定と一致しません。"
        )


def _build_hierarchical_table_group(table, region, table_texts):
    raw_records = sorted(
        build_ba_records(table_texts),
        key=lambda record: (-record["y"], record["x"]),
    )
    if len(raw_records) != 23:
        raise ExtractionError(
            f"階層型土工表の単位行が23件ではありません（{len(raw_records)}件）。"
        )
    invalid_quantities = [
        record["数量"] for record in raw_records
        if not (
            isinstance(record["数量"], Decimal)
            or record["数量"] == "-"
        )
    ]
    if invalid_quantities:
        raise ExtractionError(
            f"階層型土工表に数値化できない数量があります（{invalid_quantities}）。"
        )

    kind_headers = sorted(
        [header for header in table["headers"] if header["text"] == "種別"],
        key=lambda header: header["x"],
    )
    division_headers = sorted(
        [
            text for text in table_texts
            if text["text"] == "区分"
            and abs(text["y"] - table["y"]) <= TABLE_HEADER_Y_TOLERANCE
        ],
        key=lambda text: text["x"],
    )
    if len(kind_headers) != 2 or len(division_headers) != 2:
        raise ExtractionError("階層型土工表の種別・区分ヘッダーが不足しています。")

    categorized_records = [
        (
            raw_record,
            _hierarchical_category(
                raw_record,
                table,
                region,
                table_texts,
                kind_headers,
                division_headers,
            ),
        )
        for raw_record in raw_records
    ]
    _validate_hierarchical_source_schema(
        [
            (category, raw_record["種別"])
            for raw_record, category in categorized_records
        ]
    )

    mapped = {}

    def add_record(kind, source, quantity=None, combine=False):
        value = source["数量"] if quantity is None else quantity
        record = {
            "x": source["x"],
            "y": source["y"],
            "種別": kind,
            "単位": source["単位"],
            "数量": value,
        }
        if kind in mapped:
            if not combine:
                raise ExtractionError(f"転記項目「{kind}」が重複しています。")
            mapped[kind]["数量"] = _combine_table_quantities(
                mapped[kind]["数量"],
                value,
            )
        else:
            mapped[kind] = record

    for raw_record, category in categorized_records:
        division = raw_record["種別"]

        if category == "オープン掘削":
            add_record("オープン掘削", raw_record)
        elif (
            category in HIERARCHICAL_WIDTH_CATEGORY
            and division in HIERARCHICAL_WIDTH_SUFFIX
        ):
            base_kind = HIERARCHICAL_WIDTH_CATEGORY[category]
            add_record(
                base_kind + HIERARCHICAL_WIDTH_SUFFIX[division],
                raw_record,
                combine=base_kind == "路床盛土",
            )
        elif category in {"路肩盛土", "畦畔盛土"}:
            add_record(category, raw_record)
        elif category == "床堀り":
            add_record("床掘", raw_record)
        elif category == "埋戻し" and division in {"C", "D"}:
            add_record(f"埋戻({division})", raw_record)
        elif category == "法面整形工":
            quantity_values = _row_quantity_values(
                raw_record,
                table_texts,
            )
            if len(quantity_values) != 2:
                raise ExtractionError("法面整形工の左右数量を取得できませんでした。")
            if division.startswith("切土部"):
                prefix = "切土法面整形"
            elif division.startswith("盛土部"):
                prefix = "盛土法面整形"
            else:
                raise ExtractionError("法面整形工の切土・盛土区分を判定できません。")
            add_record(f"{prefix}(左)", raw_record, quantity_values[0])
            add_record(f"{prefix}(右)", raw_record, quantity_values[1])

    missing = [
        kind for kind in HIERARCHICAL_ALL_OUTPUT_KINDS
        if kind not in mapped
    ]
    extra = [
        kind for kind in mapped
        if kind not in HIERARCHICAL_ALL_OUTPUT_KINDS
    ]
    if missing or extra:
        raise ExtractionError(
            f"階層型土工表の転記項目が一致しません"
            f"（不足={missing}, 余分={extra}）。"
        )

    return [
        mapped[kind]
        for kind in HIERARCHICAL_ALL_OUTPUT_KINDS
    ]


def _table_row_bounds(region, x):
    tolerance = (
        region["text_height"] * LINE_CONNECTION_TOLERANCE_FACTOR
    )
    levels = _unique_axis_values(
        [
            (line["start_y"] + line["end_y"]) / 2
            for line in region["main_lines"]
            if abs(line["end_y"] - line["start_y"])
            <= LINE_AXIS_TOLERANCE
            and min(line["start_x"], line["end_x"]) - tolerance
            <= x
            <= max(line["start_x"], line["end_x"]) + tolerance
        ],
        tolerance,
    )
    if len(levels) < 2:
        raise ExtractionError("種別型土工表の行境界を取得できません。")
    return [
        (levels[index], levels[index + 1])
        for index in range(len(levels) - 1)
    ]


def _category_table_row_quantity(
    table_texts,
    row_bounds,
    quantity_x,
    division_bounds,
    expected_division,
    category,
):
    lower_y, upper_y = row_bounds
    tolerance = TOLERANCE_Y_ROW * 0.05
    quantities = [
        text for text in table_texts
        if lower_y - tolerance <= text["y"] <= upper_y + tolerance
        and abs(text["x"] - quantity_x) <= 1.5
        and _parse_table_quantity(text["text"]) is not None
    ]
    if len(quantities) != 1:
        raise ExtractionError(
            f"種別型土工表の「{category}」に対応する数量を"
            "一意に特定できません。"
        )

    division_min_x, division_max_x = division_bounds
    divisions = [
        text for text in table_texts
        if lower_y - tolerance <= text["y"] <= upper_y + tolerance
        and division_min_x <= text["x"] <= division_max_x
    ]
    actual_divisions = [text["text"] for text in divisions]
    expected_divisions = (
        [] if expected_division is None else [expected_division]
    )
    if actual_divisions != expected_divisions:
        raise ExtractionError(
            f"種別型土工表の「{category}」の区分が一致しません"
            f"（期待={expected_divisions}, 実際={actual_divisions}）。"
        )

    return _parse_table_quantity(quantities[0]["text"])


def _build_category_quantity_table_group(table, region, table_texts):
    """
    種別・区分・数量と罫線行から数量を読む。

    記号列は使用せず、左右なしの法面整形は左へ割り当てる。
    """
    headers = {
        header_text: sorted(
            [
                header for header in table["headers"]
                if header["text"] == header_text
            ],
            key=lambda header: header["x"],
        )
        for header_text in ("種別", "区分", "数量")
    }
    if any(len(values) != 2 for values in headers.values()):
        raise ExtractionError("種別型土工表のヘッダーが不足しています。")

    center_x = (
        region["main_bounds"]["min_x"]
        + region["main_bounds"]["max_x"]
    ) / 2
    mapped = {}
    for category, divisions in CATEGORY_TABLE_SOURCE_SCHEMA:
        category_texts = [
            text for text in table_texts
            if text["text"] == category
        ]
        if len(category_texts) != 1:
            raise ExtractionError(
                f"種別型土工表の種別「{category}」を"
                "一意に特定できません。"
            )

        category_text = category_texts[0]
        side_index = 0 if category_text["x"] < center_x else 1
        kind_x = headers["種別"][side_index]["x"]
        division_x = headers["区分"][side_index]["x"]
        quantity_x = headers["数量"][side_index]["x"]
        rows = _table_row_bounds(region, quantity_x)
        category_row_indexes = [
            index for index, (lower_y, upper_y) in enumerate(rows)
            if lower_y <= category_text["y"] <= upper_y
        ]
        if len(category_row_indexes) != 1:
            raise ExtractionError(
                f"種別型土工表の種別「{category}」の行を"
                "一意に特定できません。"
            )

        category_row_index = category_row_indexes[0]
        division_bounds = (
            (kind_x + division_x) / 2,
            (division_x + quantity_x) / 2,
        )
        for offset, (division, output_kind) in enumerate(divisions):
            row_index = category_row_index - offset
            if row_index < 0:
                raise ExtractionError(
                    f"種別型土工表の種別「{category}」の"
                    "数量行が不足しています。"
                )
            quantity = _category_table_row_quantity(
                table_texts,
                rows[row_index],
                quantity_x,
                division_bounds,
                division,
                category,
            )
            mapped[output_kind] = {
                "x": category_text["x"],
                "y": sum(rows[row_index]) / 2,
                "種別": output_kind,
                "単位": "",
                "数量": quantity,
            }

    for output_kind in CATEGORY_TABLE_RIGHT_OUTPUT_KINDS:
        source_kind = output_kind.replace("(右)", "(左)")
        source = mapped[source_kind]
        mapped[output_kind] = {
            "x": source["x"],
            "y": source["y"],
            "種別": output_kind,
            "単位": "",
            "数量": None,
        }

    missing = [
        kind for kind in HIERARCHICAL_OUTPUT_KINDS
        if kind not in mapped
    ]
    if missing:
        raise ExtractionError(
            f"種別型土工表の転記項目が不足しています（{missing}）。"
        )

    return [
        mapped[kind]
        for kind in HIERARCHICAL_OUTPUT_KINDS
    ]


def build_table_data_groups(all_texts, tables, table_regions):
    """各土工表の罫線領域内だけからExcel転記用データを構築する。"""
    if len(tables) != len(table_regions):
        raise ExtractionError("土工表と罫線領域の件数が一致しません。")

    groups = []
    for table, region in zip(tables, table_regions):
        table_texts = _texts_in_table_region(
            all_texts,
            table,
            region,
        )
        if table.get("format") == CATEGORY_TABLE_FORMAT:
            group = _build_category_quantity_table_group(
                table,
                region,
                table_texts,
            )
        else:
            division_headers = [
                text for text in table_texts
                if text["text"] == "区分"
                and abs(text["y"] - table["y"]) <= TABLE_HEADER_Y_TOLERANCE
            ]
        if (
            table.get("format") != CATEGORY_TABLE_FORMAT
            and division_headers
        ):
            group = _build_hierarchical_table_group(
                table,
                region,
                table_texts,
            )
        elif table.get("format") != CATEGORY_TABLE_FORMAT:
            group = _build_legacy_table_group(table_texts)
        groups.append(group)

    return groups


def assign_data_groups_to_tables(tables, data_groups):
    """各データ群を最寄りの表へ割り当て、1表1群でなければ停止する。"""
    if not tables:
        raise ExtractionError("土工表の共通ヘッダーを検出できませんでした。")

    assignments = [[] for _ in tables]
    for data_group_index, data_group in enumerate(data_groups):
        if not data_group:
            continue
        table_index = min(
            range(len(tables)),
            key=lambda i: _distance(tables[i], data_group[0]),
        )
        assignments[table_index].append(data_group_index)

    invalid = [
        (index + 1, len(group_indexes))
        for index, group_indexes in enumerate(assignments)
        if len(group_indexes) != 1
    ]
    if invalid:
        count_summary = ", ".join(
            f"{count}群の表が{table_count}件"
            for count, table_count in sorted(
                Counter(count for _, count in invalid).items()
            )
        )
        raise ExtractionError(
            f"土工表{len(tables)}件にデータ群{len(data_groups)}件を"
            f"一意に対応付けできません（{count_summary}）。"
            "現段階では誤転記を避けるため、このDXFを出力しません。"
        )

    for table_index, group_indexes in enumerate(assignments):
        data_group_index = group_indexes[0]
        data_group = data_groups[data_group_index]
        table = tables[table_index]
        if _distance(table, data_group[0]) > MAX_DATA_GROUP_TO_TABLE_DISTANCE:
            raise ExtractionError(
                "土工表とデータ群の距離が安全上限を超えています。"
                "現段階では誤転記を避けるため、このDXFを出力しません。"
            )

        nearest_group_index = _nearest_item_index(
            table,
            [group[0] for group in data_groups],
            "データ群",
        )
        if nearest_group_index != data_group_index:
            raise ExtractionError(
                "土工表とデータ群が相互に最寄りではありません。"
                "現段階では誤転記を避けるため、このDXFを出力しません。"
            )

    return [
        data_groups[group_indexes[0]]
        for group_indexes in assignments
    ]


def _nearest_item_index(origin, items, target_name):
    ranked = sorted(
        (
            (_distance(origin, item), index)
            for index, item in enumerate(items)
        ),
        key=lambda pair: pair[0],
    )
    if (
        len(ranked) >= 2
        and ranked[1][0] - ranked[0][0] < MIN_NEAREST_DISTANCE_GAP
    ):
        raise ExtractionError(
            f"最寄りの{target_name}候補が近接しているため、一意に決定できません。"
        )
    return ranked[0][1]


def resolve_table_stations(tables, stations, center_markers):
    """
    表ごとの中心測点を解決する。

    CL TEXTがある場合は「表→CL→測点」、ない場合は測点を最寄り表へ
    逆割当し、候補が1件だけの表に限って採用する。
    """
    if not tables:
        raise ExtractionError("中心測点を対応付ける土工表がありません。")
    if not stations:
        raise ExtractionError("測点TEXTを検出できませんでした。")

    if center_markers:
        center_indexes = []
        for table in tables:
            center_index = _nearest_item_index(table, center_markers, "CL")
            center = center_markers[center_index]
            if (
                _distance(table, center) > MAX_TABLE_TO_CENTER_DISTANCE
                or abs(table["y"] - center["y"]) > MAX_TABLE_TO_CENTER_Y_DISTANCE
            ):
                raise ExtractionError(
                    "土工表に対応するCLを安全な距離内で特定できませんでした。"
                )
            center_indexes.append(center_index)

        if len(set(center_indexes)) != len(center_indexes):
            raise ExtractionError("複数の土工表が同じCLに対応し、中心断面を一意に特定できません。")

        station_indexes = []
        for center_index in center_indexes:
            center = center_markers[center_index]
            station_index = _nearest_item_index(center, stations, "測点")
            if _distance(center, stations[station_index]) > MAX_CENTER_TO_STATION_DISTANCE:
                raise ExtractionError("CLに対応する測点を安全な距離内で特定できませんでした。")
            station_indexes.append(station_index)

        if len(set(station_indexes)) != len(station_indexes):
            raise ExtractionError("複数の土工表が同じ測点に対応し、中心測点を一意に特定できません。")

        return [stations[index] for index in station_indexes]

    assignments = [[] for _ in tables]
    for station_index, station in enumerate(stations):
        table_index = _nearest_item_index(station, tables, "土工表")
        assignments[table_index].append(station_index)

    invalid = [
        (index + 1, len(station_indexes))
        for index, station_indexes in enumerate(assignments)
        if len(station_indexes) != 1
    ]
    if invalid:
        count_summary = ", ".join(
            f"{count}測点の表が{table_count}件"
            for count, table_count in sorted(
                Counter(count for _, count in invalid).items()
            )
        )
        raise ExtractionError(
            "CL TEXTがなく、各土工表に対応する測点を1件に絞れません"
            f"（{count_summary}）。"
        )

    for table_index, station_indexes in enumerate(assignments):
        station_index = station_indexes[0]
        table = tables[table_index]
        station = stations[station_index]
        if (
            _distance(table, station) > MAX_TABLE_TO_STATION_DISTANCE
            or abs(table["y"] - station.get("anchor_y", station["y"]))
            > MAX_TABLE_TO_STATION_Y_DISTANCE
        ):
            raise ExtractionError(
                "CL TEXTがない土工表と測点の距離が安全上限を超えています。"
            )

        nearest_station_index = _nearest_item_index(table, stations, "測点")
        if nearest_station_index != station_index:
            raise ExtractionError(
                "CL TEXTがない土工表と測点が相互に最寄りではありません。"
            )

    return [
        stations[station_indexes[0]]
        for station_indexes in assignments
    ]


def _select_mainline_tables(
    tables,
    table_regions,
    all_texts,
):
    """明示された路線名がある場合は、本線の土工表だけを残す。"""
    route_names = []
    for table, region in zip(tables, table_regions):
        main_bounds = region["main_bounds"]
        candidates = [
            text for text in all_texts
            if text["layer"] == table["layer"]
            and ROUTE_NAME_PATTERN.fullmatch(text["text"])
            and main_bounds["min_x"] - region["text_height"]
            <= text["x"]
            <= main_bounds["max_x"] + region["text_height"]
            and main_bounds["max_y"] < text["y"]
            <= main_bounds["max_y"] + region["text_height"] * 2
        ]
        if len(candidates) > 1:
            raise ExtractionError(
                "土工表の路線名を一意に特定できません。"
            )
        route_names.append(
            None if not candidates else candidates[0]["text"]
        )

    explicit_route_count = sum(
        route_name is not None
        for route_name in route_names
    )
    if explicit_route_count == 0:
        return tables, table_regions, ["本線"] * len(tables)
    if explicit_route_count != len(tables):
        raise ExtractionError(
            "路線名のある土工表とない土工表が混在しています。"
        )

    mainline_indexes = [
        index for index, route_name in enumerate(route_names)
        if route_name == "本線"
    ]
    if not mainline_indexes:
        raise ExtractionError("本線の土工表を検出できませんでした。")

    return (
        [tables[index] for index in mainline_indexes],
        [table_regions[index] for index in mainline_indexes],
        ["本線"] * len(mainline_indexes),
    )


def _minimum_cost_unique_assignment(origins, candidates):
    """矩形ハンガリアン法で各表へ異なる測点を最小距離で割り当てる。"""
    origin_count = len(origins)
    candidate_count = len(candidates)
    if origin_count > candidate_count:
        raise ExtractionError("土工表より測点候補が少ないため対応付けできません。")

    row_potentials = [0.0] * (origin_count + 1)
    column_potentials = [0.0] * (candidate_count + 1)
    column_rows = [0] * (candidate_count + 1)
    previous_columns = [0] * (candidate_count + 1)

    for row in range(1, origin_count + 1):
        column_rows[0] = row
        current_column = 0
        minimum_values = [float("inf")] * (candidate_count + 1)
        used = [False] * (candidate_count + 1)

        while True:
            used[current_column] = True
            current_row = column_rows[current_column]
            delta = float("inf")
            next_column = 0
            for column in range(1, candidate_count + 1):
                if used[column]:
                    continue
                reduced_cost = (
                    _distance(
                        origins[current_row - 1],
                        candidates[column - 1],
                    )
                    - row_potentials[current_row]
                    - column_potentials[column]
                )
                if reduced_cost < minimum_values[column]:
                    minimum_values[column] = reduced_cost
                    previous_columns[column] = current_column
                if minimum_values[column] < delta:
                    delta = minimum_values[column]
                    next_column = column

            for column in range(candidate_count + 1):
                if used[column]:
                    row_potentials[column_rows[column]] += delta
                    column_potentials[column] -= delta
                else:
                    minimum_values[column] -= delta

            current_column = next_column
            if column_rows[current_column] == 0:
                break

        while True:
            previous_column = previous_columns[current_column]
            column_rows[current_column] = column_rows[previous_column]
            current_column = previous_column
            if current_column == 0:
                break

    assignments = [None] * origin_count
    for column in range(1, candidate_count + 1):
        row = column_rows[column]
        if row != 0:
            assignments[row - 1] = column - 1

    if any(index is None for index in assignments):
        raise ExtractionError("土工表と測点の最小距離対応付けに失敗しました。")
    return assignments


def _resolve_category_table_stations(tables, stations, alignments):
    grouped_indexes = {}
    for index, alignment in enumerate(alignments):
        grouped_indexes.setdefault(alignment, []).append(index)

    resolved = [None] * len(tables)
    for alignment, table_indexes in grouped_indexes.items():
        mainline = alignment == "本線"
        candidates = [
            station for station in stations
            if (
                int(station["text"]) >= CATEGORY_TABLE_MAINLINE_STATION_MIN
            ) == mainline
        ]
        if not candidates:
            raise ExtractionError(
                f"路線「{alignment}」の測点候補がありません。"
            )

        alignment_tables = [
            tables[index]
            for index in table_indexes
        ]
        candidate_indexes = _minimum_cost_unique_assignment(
            alignment_tables,
            candidates,
        )
        alignment_stations = [
            candidates[index]
            for index in candidate_indexes
        ]

        for table, station in zip(
            alignment_tables,
            alignment_stations,
        ):
            if (
                _distance(table, station) > MAX_TABLE_TO_STATION_DISTANCE
                or abs(
                    table["y"]
                    - station.get("anchor_y", station["y"])
                ) > MAX_TABLE_TO_STATION_Y_DISTANCE
            ):
                raise ExtractionError(
                    f"路線「{alignment}」の土工表と測点の距離が"
                    "安全上限を超えています。"
                )

        ordered_numbers = sorted(
            int(station["text"])
            for station in alignment_stations
        )
        expected_numbers = list(
            range(ordered_numbers[0], ordered_numbers[-1] + 1)
        )
        if ordered_numbers != expected_numbers:
            raise ExtractionError(
                f"路線「{alignment}」の対応測点が連続していません。"
            )

        for table_index, station in zip(
            table_indexes,
            alignment_stations,
        ):
            resolved[table_index] = station

    return resolved


def _records_for_tables(
    table_stations,
    table_data_groups,
    alignments=None,
):
    out_data = [
        {
            "測点": Decimal(station["text"]),
            "追加距離": station["plus"],
            "データ": data_group,
            **(
                {"路線": alignment}
                if alignment is not None
                else {}
            ),
        }
        for station, data_group, alignment in zip(
            table_stations,
            table_data_groups,
            alignments or [None] * len(table_stations),
        )
    ]
    out_data.sort(key=sort_key)
    return out_data


def extract_output_record_sets(all_texts, all_lines):
    """DXFから、同じ測点系列でExcel化できるレコード群を抽出する。"""
    tables = detect_earthwork_tables(all_texts)
    table_regions = detect_table_regions(
        tables,
        all_lines,
        all_texts,
    )
    tables, table_regions, alignments = _select_mainline_tables(
        tables,
        table_regions,
        all_texts,
    )
    data_groups = build_table_data_groups(
        all_texts,
        tables,
        table_regions,
    )

    stations = extract_no_texts(all_texts)
    formats = {
        table.get("format", "unit")
        for table in tables
    }
    if formats == {CATEGORY_TABLE_FORMAT}:
        table_stations = _resolve_category_table_stations(
            tables,
            stations,
            alignments,
        )
        return {
            "本線": _records_for_tables(
                table_stations,
                data_groups,
                alignments,
            )
        }

    if CATEGORY_TABLE_FORMAT in formats:
        raise ExtractionError(
            "異なる形式の土工表が同じDXFに混在しています。"
        )

    table_data_groups = assign_data_groups_to_tables(tables, data_groups)
    center_markers = extract_center_markers(all_texts)
    table_stations = resolve_table_stations(
        tables,
        stations,
        center_markers,
    )
    return {
        None: _records_for_tables(
            table_stations,
            table_data_groups,
        )
    }


def extract_output_records(all_texts, all_lines):
    """DXFのTEXT・LINEから、Excel転記前の全レコードを抽出する。"""
    record_sets = extract_output_record_sets(
        all_texts,
        all_lines,
    )
    return [
        record
        for records in record_sets.values()
        for record in records
    ]


# ==================================================
# 可変件数のExcel帳票
# ==================================================

AVERAGE_FORMULA_PATTERN = re.compile(
    r"\((\$?[A-Z]{1,3}\$?\d+)\+(\$?[A-Z]{1,3}\$?\d+)\)/2"
)
CELL_ROW_REFERENCE_PATTERN = re.compile(
    r"(\$?[A-Z]{1,3}\$?)(\d+)"
)


def infer_station_interval(out_data):
    """測点列が厳密に増加する最小の測点間隔を選ぶ。"""
    if not out_data:
        raise ExtractionError("測点がありません。")

    for interval in STATION_INTERVAL_CANDIDATES:
        positions = [
            (
                item["測点"] * interval
                + (
                    Decimal("0")
                    if item["追加距離"] is None
                    else item["追加距離"]
                )
            )
            for item in out_data
        ]
        if all(
            current > previous
            for previous, current in zip(positions, positions[1:])
        ):
            return interval

    raise ExtractionError(
        "測点間隔を20mまたは100mのいずれにも決定できませんでした。"
    )


def _make_formula_dash_safe(formula):
    return AVERAGE_FORMULA_PATTERN.sub(
        lambda match: (
            f"SUM({match.group(1)}:{match.group(2)})/2"
        ),
        formula,
    )


def _remap_formula_rows(formula, row_mapping):
    def replace(match):
        row = int(match.group(2))
        mapped_row = row_mapping.get(row, row)
        return f"{match.group(1)}{mapped_row}"

    return CELL_ROW_REFERENCE_PATTERN.sub(replace, formula)


def _snapshot_row(ws, row):
    dimension = ws.row_dimensions[row]
    return {
        "row": row,
        "height": dimension.height,
        "hidden": dimension.hidden,
        "outline_level": dimension.outlineLevel,
        "collapsed": dimension.collapsed,
        "cells": [
            {
                "coordinate": cell.coordinate,
                "value": cell.value,
                "style": copy(cell._style),
                "comment": copy(cell.comment),
            }
            for cell in ws[row]
        ],
    }


def _apply_row_snapshot(
    ws,
    snapshot,
    destination_row,
    *,
    translate=False,
    row_mapping=None,
    dash_safe=False,
):
    for column, source in enumerate(snapshot["cells"], start=1):
        target = ws.cell(destination_row, column)
        target._style = copy(source["style"])
        target.comment = copy(source["comment"])

        value = source["value"]
        if isinstance(value, str) and value.startswith("="):
            if row_mapping is not None:
                value = _remap_formula_rows(value, row_mapping)
            elif translate:
                value = Translator(
                    value,
                    origin=source["coordinate"],
                ).translate_formula(target.coordinate)
            if dash_safe:
                value = _make_formula_dash_safe(value)
        target.value = value

    dimension = ws.row_dimensions[destination_row]
    dimension.height = snapshot["height"]
    dimension.hidden = snapshot["hidden"]
    dimension.outlineLevel = snapshot["outline_level"]
    dimension.collapsed = snapshot["collapsed"]


def _update_detail_defined_names(
    ws,
    data_last_row,
    report_last_row,
):
    data_names = {
        "_1B",
        "_1F",
        "_2B",
        "_2F",
        "_3F",
        "_4F",
        "_Fill",
    }
    report_names = {
        "_1P",
        "AREA",
        "AREA1",
        "Print_Area_MI",
    }

    for name, defined_name in ws.defined_names.items():
        if name in data_names:
            end_row = data_last_row
        elif name in report_names:
            end_row = report_last_row
        else:
            continue
        defined_name.attr_text = re.sub(
            r"\$\d+$",
            f"${end_row}",
            defined_name.attr_text,
        )


def _configure_detail_sheet(ws, station_count, *, slope=False):
    data_last_row = ROW_START + station_count
    if slope:
        template_data_last = 39
        template_total_rows = (40, 41, 42)
        subtotal_row = max(40, data_last_row + 1)
        total_row = subtotal_row + 1
        side_total_row = total_row + 1
        target_total_rows = (
            subtotal_row,
            total_row,
            side_total_row,
        )
        row_mapping = {
            template_data_last: data_last_row,
            40: subtotal_row,
            41: total_row,
            42: side_total_row,
        }
        print_last_row = side_total_row
        report_last_row = total_row
    else:
        template_data_last = 40
        template_total_rows = (41, 42)
        subtotal_row = max(41, data_last_row + 1)
        total_row = subtotal_row + 1
        target_total_rows = (subtotal_row, total_row)
        row_mapping = {
            template_data_last: data_last_row,
            41: subtotal_row,
            42: total_row,
        }
        side_total_row = None
        print_last_row = total_row
        report_last_row = total_row

    detail_template = _snapshot_row(ws, 7)
    total_templates = [
        _snapshot_row(ws, row)
        for row in template_total_rows
    ]

    for row in range(7, data_last_row + 1):
        _apply_row_snapshot(
            ws,
            detail_template,
            row,
            translate=True,
            dash_safe=True,
        )
        ws.cell(row, 1).value = row - ROW_START

    for snapshot, target_row in zip(
        total_templates,
        target_total_rows,
    ):
        _apply_row_snapshot(
            ws,
            snapshot,
            target_row,
            row_mapping=row_mapping,
            dash_safe=True,
        )

    end_column = "T" if ws.title == "路体" else "R"
    ws.print_area = f"B1:{end_column}{print_last_row}"
    ws.print_title_rows = "1:5"
    _update_detail_defined_names(
        ws,
        data_last_row,
        report_last_row,
    )

    for page_break_row in range(40, data_last_row, 35):
        ws.row_breaks.append(Break(id=page_break_row))

    return {
        "data_last_row": data_last_row,
        "subtotal_row": subtotal_row,
        "total_row": total_row,
        "side_total_row": side_total_row,
    }


def _set_input_formulas(ws, out_data):
    station_count = len(out_data)
    input_last_row = ROW_START + station_count - 1
    if input_last_row > ws.max_row:
        raise ExtractionError(
            f"入力シートの上限を超えています（{station_count}件）。"
        )

    interval = infer_station_interval(out_data)
    ws.cell(row=4, column=5).value = int(interval)

    for row in range(ROW_START, input_last_row + 1):
        if row > ROW_START:
            ws.cell(row=row, column=6).value = (
                f'=IF(ISBLANK(C{row}),"",'
                f"(C{row}-C{row - 1})*$E$4"
                f"+(E{row}-E{row - 1}))"
            )
        ws[f"AE{row}"] = f"=G{row}"
        ws[f"AF{row}"] = f"=SUM(K{row}:X{row})"
        ws[f"AG{row}"] = f"=SUM(Y{row}:Z{row})"
        ws[f"AH{row}"] = f"=SUM(AA{row}:AB{row})"

    return input_last_row


def _quantity_by_kind(record):
    return {
        data["種別"]: data["数量"]
        for data in record["データ"]
    }


def _populate_work_earthwork_sheet(ws, out_data, layout):
    ws["M3"] = "埋戻(C)"

    for index, record in enumerate(out_data):
        row = ROW_START + 1 + index
        quantities = _quantity_by_kind(record)
        if set(HIERARCHICAL_WORK_OUTPUT_KINDS) <= quantities.keys():
            work_quantities = {
                kind: quantities[kind]
                for kind in HIERARCHICAL_WORK_OUTPUT_KINDS
            }
        elif {"床掘", "埋戻"} <= quantities.keys():
            work_quantities = {
                "床掘": quantities["床掘"],
                "埋戻(C)": "-",
                "埋戻(D)": quantities["埋戻"],
            }
        elif not (
            {
                "床掘",
                "埋戻",
                "埋戻(C)",
                "埋戻(D)",
            }
            & quantities.keys()
        ):
            work_quantities = {
                "床掘": None,
                "埋戻(C)": None,
                "埋戻(D)": None,
            }
        else:
            raise ExtractionError(
                "作業土工の床掘・埋戻数量を特定できませんでした。"
            )

        for column, kind in (
            (7, "床掘"),
            (10, "埋戻(D)"),
            (13, "埋戻(C)"),
        ):
            cell = ws.cell(row=row, column=column)
            cell.value = work_quantities[kind]
            cell.alignment = Alignment(horizontal="right")

        for column in (16, 17, 18):
            ws.cell(row=row, column=column).value = None

        if row == ROW_START + 1:
            for column in (8, 9, 11, 12, 14, 15):
                ws.cell(row=row, column=column).value = None
            continue

        previous_row = row - 1
        for area_column, average_column, volume_column in (
            ("G", "H", "I"),
            ("J", "K", "L"),
            ("M", "N", "O"),
        ):
            ws[f"{average_column}{row}"] = (
                f'=IF(ISBLANK($C{row}),"",'
                f"SUM({area_column}{previous_row}:"
                f"{area_column}{row})/2)"
            )
            ws[f"{volume_column}{row}"] = (
                f'=ROUND(IF(ISBLANK($C{row}),"",'
                f"$F{row}*{average_column}{row}),1)"
            )

    subtotal_row = layout["subtotal_row"]
    total_row = layout["total_row"]
    for column in ("I", "L", "O"):
        ws[f"{column}{subtotal_row}"] = (
            f"=SUM({column}7:{column}{layout['data_last_row']})"
        )
        ws[f"{column}{total_row}"] = f"={column}{subtotal_row}"


def _add_road_shoulder_to_roadbed(ws, layout):
    ws["P3"] = '=IF(入力!T2="","",入力!T2)'
    ws["P6"] = "=入力!T5"

    for row in range(7, layout["data_last_row"] + 1):
        input_row = row - 1
        ws[f"P{row}"] = f"=入力!T{input_row}"
        ws[f"Q{row}"] = (
            f'=IF(ISBLANK($C{row}),"",'
            f"SUM(P{row - 1}:P{row})/2)"
        )
        ws[f"R{row}"] = (
            f'=ROUND(IF(ISBLANK($C{row}),"",'
            f"$F{row}*Q{row}),1)"
        )

    subtotal_row = layout["subtotal_row"]
    total_row = layout["total_row"]
    ws[f"R{subtotal_row}"] = (
        f"=SUM(R7:R{layout['data_last_row']})"
    )
    ws[f"R{total_row}"] = f"=R{subtotal_row}"


def _repair_check_sheet(ws, input_last_row, station_count):
    ws["C4"] = (
        f"=INDEX(入力!$C$5:$C${input_last_row},$C$2)"
    )
    ws["E4"] = (
        f"=INDEX(入力!$E$5:$E${input_last_row},$C$2)"
    )

    for row in range(5, 33):
        ws[f"C{row}"] = None
        ws[f"D{row}"] = None

    for cell, value in {
        "B5": "掘削",
        "B9": "盛土①",
        "B13": "盛土②",
        "B17": "盛土③",
        "B21": "法面",
        "B25": None,
        "B29": None,
    }.items():
        ws[cell] = value

    row_to_input_column = {
        5: "G",
        9: "K",
        10: "N",
        11: "O",
        13: "Q",
        14: "R",
        15: "S",
        16: "T",
        17: "U",
        18: "V",
        19: "W",
        20: "X",
        21: "Y",
        22: "Z",
        23: "AA",
        24: "AB",
    }
    for row, column in row_to_input_column.items():
        ws[f"C{row}"] = (
            f'=IF(入力!{column}$2="","",入力!{column}$2)'
        )
        ws[f"D{row}"] = (
            f"=INDEX(入力!${column}$5:"
            f"${column}${input_last_row},$C$2)"
        )

    ws.data_validations.dataValidation = []
    validation = DataValidation(
        type="whole",
        operator="between",
        formula1="1",
        formula2=str(station_count),
        allow_blank=False,
    )
    validation.error = f"1から{station_count}までを入力してください。"
    validation.errorTitle = "セル番号が範囲外です"
    validation.prompt = (
        f"確認する測点のセル番号（1～{station_count}）"
    )
    validation.promptTitle = "セル番号"
    validation.showErrorMessage = True
    validation.showInputMessage = True
    ws.add_data_validation(validation)
    validation.add(ws["C2"])


def _update_downstream_formulas(wb, layouts):
    standard_total = layouts["掘削"]["total_row"]
    slope_side_total = layouts["法面整形"]["side_total_row"]
    work_total = layouts["作業土工1"]["total_row"]

    summary = wb["集計表"]
    for cell, formula in {
        "F5": f"=掘削!I{standard_total}",
        "F7": f"=路体!I{layouts['路体']['total_row']}",
        "F8": f"=路体!N{layouts['路体']['total_row']}",
        "F9": f"=路体!Q{layouts['路体']['total_row']}",
        "F10": f"=路床!I{layouts['路床']['total_row']}",
        "F11": f"=路床!L{layouts['路床']['total_row']}",
        "F12": f"=路床!O{layouts['路床']['total_row']}",
        "F13": f"=路体外!I{layouts['路体外']['total_row']}",
        "F14": f"=路体外!L{layouts['路体外']['total_row']}",
        "F15": f"=路体外!O{layouts['路体外']['total_row']}",
        "F16": f"=路体外!R{layouts['路体外']['total_row']}",
        "F20": f"=法面整形!L{slope_side_total}",
        "F21": f"=法面整形!R{slope_side_total}",
    }.items():
        summary[cell] = formula
    summary["C17"] = "路肩盛土"
    summary["E17"] = "m3"
    summary["F17"] = f"=路床!R{layouts['路床']['total_row']}"
    summary["G17"] = "BA3"

    surplus = wb["残土"]
    for cell, formula in {
        "E5": f"=掘削!I{standard_total}",
        "E9": f"=路体!I{layouts['路体']['total_row']}",
        "E10": f"=路体!N{layouts['路体']['total_row']}",
        "E11": f"=路体!Q{layouts['路体']['total_row']}",
        "E13": f"=路床!I{layouts['路床']['total_row']}",
        "E14": f"=路床!L{layouts['路床']['total_row']}",
        "E15": f"=路床!O{layouts['路床']['total_row']}",
        "E17": f"=路体外!I{layouts['路体外']['total_row']}",
        "E18": f"=路体外!L{layouts['路体外']['total_row']}",
        "E19": f"=路体外!O{layouts['路体外']['total_row']}",
        "E21": f"=路体外!R{layouts['路体外']['total_row']}",
    }.items():
        surplus[cell] = formula
    surplus["B16"] = "路肩盛土"
    surplus["D16"] = "m3"
    surplus["E16"] = f"=路床!R{layouts['路床']['total_row']}"
    surplus["F16"] = Decimal("0.9")
    surplus["G16"] = "=ROUND(E16/F16,1)"

    work_summary = wb["作業土工2"]
    work_summary["G5"] = f"=作業土工1!I{work_total}"
    work_summary["H5"] = f"=作業土工1!O{work_total}"
    work_summary["I5"] = f"=作業土工1!L{work_total}"


def _prepare_dynamic_workbook(wb, out_data):
    station_count = len(out_data)
    input_sheet = wb[SHEET_NAME]
    input_last_row = _set_input_formulas(input_sheet, out_data)

    layouts = {}
    for sheet_name in (
        "掘削",
        "路体",
        "路床",
        "路体外",
        "作業土工1",
    ):
        layouts[sheet_name] = _configure_detail_sheet(
            wb[sheet_name],
            station_count,
        )
    layouts["法面整形"] = _configure_detail_sheet(
        wb["法面整形"],
        station_count,
        slope=True,
    )

    _add_road_shoulder_to_roadbed(
        wb["路床"],
        layouts["路床"],
    )
    _populate_work_earthwork_sheet(
        wb["作業土工1"],
        out_data,
        layouts["作業土工1"],
    )
    _update_downstream_formulas(wb, layouts)
    _repair_check_sheet(
        wb["チェック"],
        input_last_row,
        station_count,
    )

    wb.calculation.calcMode = "auto"
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True


# ==================================================
# メイン
# ==================================================

def write_output_workbook(template_path, output_path, out_data):
    wb = load_workbook(template_path)
    try:
        is_hierarchical = bool(out_data) and all(
            set(HIERARCHICAL_WORK_OUTPUT_KINDS)
            <= {
                data["種別"]
                for data in record["データ"]
            }
            for record in out_data
        )
        is_category_quantity = bool(out_data) and all(
            record.get("路線")
            and [data["種別"] for data in record["データ"]]
            == list(HIERARCHICAL_OUTPUT_KINDS)
            for record in out_data
        )
        if (
            len(out_data) > VERIFIED_LEGACY_RECORD_COUNT
            or is_hierarchical
            or is_category_quantity
        ):
            _prepare_dynamic_workbook(wb, out_data)

        ws = wb[SHEET_NAME]

        for r, item in enumerate(out_data):
            row = ROW_START + r
            ws.cell(row=row, column=3).value = item["測点"]
            ws.cell(row=row, column=5).value = item["追加距離"]

            for data in item["データ"]:
                for col in range(1, ws.max_column + 1):
                    header = ws.cell(row=2, column=col)
                    if (
                        header.value is not None
                        and normalize_text(header.value) == normalize_text(data["種別"])
                    ):
                        cell = ws.cell(row=row, column=col)
                        cell.value = data["数量"]
                        cell.alignment = Alignment(horizontal="right")

        wb.save(output_path)
    finally:
        wb.close()


def proc(dxf_files):
    base = get_base_path()
    input_dir = os.path.join(base, "input")
    output_dir = os.path.join(base, "output")
    template_path = os.path.join(base, "template", TEMPLATE_NAME)

    os.makedirs(output_dir, exist_ok=True)
    timestamp = datetime.now().strftime("%y%m%d_%H%M%S")

    input_subfolder  = os.path.join(input_dir,  f"used_{timestamp}")
    os.makedirs(input_subfolder,  exist_ok=True)

    output_subfolder = os.path.join(output_dir, f"exec_{timestamp}")
    os.makedirs(output_subfolder, exist_ok=True)

    failures = []
    for dxf_file in glob.glob(os.path.join(input_dir, "*.dxf")):
        print(dxf_file)
        try:
            doc = ezdxf.readfile(dxf_file)
            modelspace = doc.modelspace()
            all_texts = collect_all_texts(modelspace)
            all_lines = collect_all_lines(modelspace)
            record_sets = extract_output_record_sets(
                all_texts,
                all_lines,
            )

            name = (
                TEMPLATE_BASE
                + "_"
                + os.path.splitext(os.path.basename(dxf_file))[0]
            )
            for alignment, out_data in record_sets.items():
                alignment_suffix = (
                    ""
                    if alignment is None
                    else "_" + alignment
                )
                output_path = os.path.join(
                    output_subfolder,
                    name + alignment_suffix + ".xlsx",
                )
                write_output_workbook(
                    template_path,
                    output_path,
                    out_data,
                )
            shutil.move(dxf_file, input_subfolder)
        except Exception as exc:
            failures.append(f"{os.path.basename(dxf_file)}: {exc}")
            continue

        dxf_files.append(os.path.basename(dxf_file))

    return dxf_files, failures


if __name__ == "__main__":
    print("oudan_red_yellow_checker [Version 1.0.0]")

    # Tkinterの初期化
    root = tk.Tk()
    root.withdraw()
    root.attributes('-topmost', True)

    if messagebox.askokcancel("処理確認", "DXFファイルの処理を開始しますか？"):
        dxf_files, failures = proc([])
        if failures:
            result_lines = []
            if dxf_files:
                result_lines.append("処理完了:\n" + "\n".join(dxf_files))
            result_lines.append("未処理:\n" + "\n".join(failures))
            messagebox.showwarning("処理結果", "\n\n".join(result_lines))
        else:
            messagebox.showinfo("処理完了", "\n".join(dxf_files))
